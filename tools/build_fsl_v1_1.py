from pathlib import Path
import os
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(os.environ.get('FSL_OUT', 'build/fsl-v1.1'))
OUT.mkdir(parents=True, exist_ok=True)
TPL = 'Four-Square Template v1.1'
NAVY='17365D'; BLUE='2F5597'; GRAY='666666'; DARK='222222'; WHITE='FFFFFF'; LIGHT='F3F6FA'; LINE='C8D0DA'

def run(p, text, size=7.0, bold=False, color=DARK, italic=False):
    r=p.add_run(text); r.font.name='Liberation Sans'; r.font.size=Pt(size); r.bold=bold; r.italic=italic; r.font.color.rgb=RGBColor.from_string(color); return r

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); sh=tcPr.find(qn('w:shd'))
    if sh is None: sh=OxmlElement('w:shd'); tcPr.append(sh)
    sh.set(qn('w:fill'), fill)

def margin(cell, n=45):
    tcPr=cell._tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None: tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for k in ('top','start','bottom','end'):
        e=tcMar.find(qn('w:'+k))
        if e is None: e=OxmlElement('w:'+k); tcMar.append(e)
        e.set(qn('w:w'), str(n)); e.set(qn('w:type'),'dxa')

def borders(table, color=LINE, sz='4'):
    pr=table._tbl.tblPr; b=pr.first_child_found_in('w:tblBorders')
    if b is None: b=OxmlElement('w:tblBorders'); pr.append(b)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=b.find(qn('w:'+edge))
        if e is None: e=OxmlElement('w:'+edge); b.append(e)
        e.set(qn('w:val'),'single'); e.set(qn('w:sz'),sz); e.set(qn('w:color'),color)

def cell(cell, text, size=6.6, bold=False, color=DARK, center=False):
    margin(cell); p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
    if center: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run(p,text,size,bold,color)

def base(product, project_version, source):
    d=Document(); s=d.sections[0]
    s.page_width=Mm(210); s.page_height=Mm(297); s.top_margin=Mm(16); s.bottom_margin=Mm(14); s.left_margin=Mm(13); s.right_margin=Mm(13); s.header_distance=Mm(5); s.footer_distance=Mm(5)
    st=d.styles['Normal']; st.font.name='Liberation Sans'; st.font.size=Pt(7.2); st.paragraph_format.space_after=Pt(1); st.paragraph_format.line_spacing=.95
    h=s.header; t=h.add_table(rows=1,cols=2,width=Mm(184)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    p=t.cell(0,0).paragraphs[0]; run(p,product,13.5,True,NAVY); run(p,'  /  FOUR-SQUARE LETTERHEAD',7.8,True,BLUE)
    p=t.cell(0,1).paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; run(p,project_version,6.8,True,NAVY); run(p,'\n'+TPL,6.1,True,GRAY); run(p,'\n'+source,5.7,False,GRAY)
    borders(t,NAVY,'7')
    f=s.footer; ft=f.add_table(rows=1,cols=3,width=Mm(184)); ft.alignment=WD_TABLE_ALIGNMENT.CENTER
    p=ft.cell(0,0).paragraphs[0]; run(p,TPL,5.8,True,NAVY); run(p,'\nSPEC + MEMORY fixed',5.3,False,GRAY)
    p=ft.cell(0,1).paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; run(p,'main1 = immutable root\n2 project lanes chosen once',5.3,False,GRAY)
    p=ft.cell(0,2).paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; run(p,'main = accepted integration\n1–3 pages max',5.3,False,GRAY)
    borders(ft,NAVY,'4')
    return d

def heading(d, text):
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(.7); run(p,text.upper(),7.8,True,NAVY)

def lineage(d, a, b):
    heading(d,'Four-Square release lineage')
    t=d.add_table(rows=3,cols=5); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.cell(0,0).merge(t.cell(0,4)); shade(t.cell(0,0),'E9EEF4'); cell(t.cell(0,0),'main1 = repository first commit / immutable root reference',6.7,True,NAVY,True)
    for ri,row in enumerate([('V1','spec-v1',a+'-v1',b+'-v1','memory-v1'),('V2','spec-v2',a+'-v2',b+'-v2','memory-v2')],1):
        for ci,val in enumerate(row):
            c=t.cell(ri,ci)
            if ci==0: shade(c,NAVY); cell(c,val,6.8,True,WHITE,True)
            else: shade(c,LIGHT if ri==1 else 'FFFFFF'); cell(c,val,6.6,ri==1,NAVY if ri==1 else DARK,True)
    borders(t)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; run(p,'SPEC + MEMORY are fixed. The two middle lanes are project-defined and remain stable across versions.',6.2,True,GRAY)

def identity(d, mission, principle, boundary):
    t=d.add_table(rows=1,cols=3); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,(h,v) in enumerate([('MISSION',mission),('CORE PRINCIPLE',principle),('BOUNDARY',boundary)]):
        shade(t.cell(0,i),LIGHT); p=t.cell(0,i).paragraphs[0]; run(p,h,6.8,True,NAVY); p=t.cell(0,i).add_paragraph(); run(p,v,6.15)
    borders(t)

def table3(d, title, rows, headers=('ITEM','PROVEN SOURCE / MECHANISM','REQUIRED OUTCOME')):
    heading(d,title); t=d.add_table(rows=1,cols=3); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers): shade(t.cell(0,i),NAVY); cell(t.cell(0,i),h,6.1,True,WHITE)
    for a,b,c in rows:
        r=t.add_row().cells; cell(r[0],a,6.15,True,NAVY); cell(r[1],b,6.05); cell(r[2],c,6.05)
    borders(t)

def lanes(d,a,b,desc_a,desc_b):
    heading(d,'Four-Square responsibilities'); t=d.add_table(rows=1,cols=4); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    vals=[('SPEC','spec-vN','Contract, workflows, invariants and acceptance criteria.'),(a.upper(),a+'-vN',desc_a),(b.upper(),b+'-vN',desc_b),('MEMORY','memory-vN','Human + AI glossary, mnemonics, examples and recall maps.')]
    for i,(h,br,desc) in enumerate(vals):
        shade(t.cell(0,i),('EEF3F8','EAF2F8','EEF6F4','F3EFF8')[i]); p=t.cell(0,i).paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; run(p,h,6.8,True,NAVY); p=t.cell(0,i).add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; run(p,br,6.0,True,BLUE); p=t.cell(0,i).add_paragraph(); run(p,desc,5.9)
    borders(t)

def banner(d):
    t=d.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; shade(t.cell(0,0),NAVY); p=t.cell(0,0).paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; run(p,'ONE VERSION · FOUR SQUARES · ALL GREEN',7.0,True,WHITE); p=t.cell(0,0).add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; run(p,'Draft ≠ Tested ≠ Accepted ≠ Released. Conformance is the arbiter.',5.9,False,WHITE); borders(t,NAVY,'5')

def page2(d,a,b):
    d.add_page_break(); heading(d,'Capability / evidence detail')
    identity(d,'Use only when Page 1 would become crowded.','Preserve the same visual identity and template version.','Add detail; do not introduce a second contract.')
    table3(d,'Lane selection rules',[(a+'-vN','Primary project-specific concern','Narrow, durable branch token.'),(b+'-vN','Second project-specific concern','Independent concern that can validate itself.'),('ai-vN','Umbrella for agents + skills + plugins','Use only when multiple AI-facing artifact classes belong together.'),('tools / code / docs / utils','Examples, not fixed names','Choose the narrowest stable project term.')])
    table3(d,'Evidence / compatibility', [('SPEC','schema / lint / contract consistency','Normative behavior coherent.'),(a.upper(),'project-specific tests/build checks','Primary deliverable proves itself.'),(b.upper(),'project-specific coverage/validation','Second lane proves itself.'),('MEMORY','glossary / link / recall-map checks','Recall aids stay aligned and non-normative.')],headers=('LANE','VALIDATION','RESULT'))
    banner(d)

def page3(d,a,b):
    d.add_page_break(); heading(d,'Release / compatibility ledger')
    table3(d,'Release certification', [('SPEC','Version aligned','Accepted'),(a.upper(),'Version aligned','Accepted'),(b.upper(),'Version aligned','Accepted'),('MEMORY','Version aligned','Accepted')],headers=('LANE','RELEASE CONDITION','STATE'))
    table3(d,'Compatibility / migration', [('[Platform / version]','[Supported path]','[Evidence / limit]'),('[Platform / version]','[Supported path]','[Evidence / limit]'),('[Platform / version]','[Supported path]','[Evidence / limit]'),('[Platform / version]','[Supported path]','[Evidence / limit]')],headers=('SCOPE','PATH','EVIDENCE / LIMIT'))
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; run(p,'State what this release proves, and what it does not prove.',6.7,True,NAVY)
    banner(d)

def make_template(pages):
    d=base('[PROJECT / PRODUCT]','[PROJECT VERSION]','[SOURCE OF TRUTH]')
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; run(p,'[SHORT PRODUCT IDENTITY / TAGLINE]',8.2,True,NAVY)
    lineage(d,'[lane-a]','[lane-b]'); identity(d,'[ONE-SENTENCE MISSION]','[CORE PRINCIPLE / INVARIANT]','[WHAT THE PROJECT OWNS / DOES NOT OWN]')
    table3(d,'Canonical workflows',[('[WORKFLOW 1]','[PROVEN ASSET]','[OUTCOME]'),('[WORKFLOW 2]','[PROVEN ASSET]','[OUTCOME]'),('[WORKFLOW 3]','[PROVEN ASSET]','[OUTCOME]'),('[WORKFLOW 4]','[PROVEN ASSET]','[OUTCOME]'),('[WORKFLOW 5]','[PROVEN ASSET]','[OUTCOME]')])
    lanes(d,'[lane-a]','[lane-b]','[PRIMARY PROJECT DELIVERABLE / CONCERN]','[SECOND PROJECT DELIVERABLE / CONCERN]'); banner(d)
    if pages>=2: page2(d,'[lane-a]','[lane-b]')
    if pages>=3: page3(d,'[lane-a]','[lane-b]')
    d.save(OUT/f'FSL-v1.1-Template-{pages}P.docx')

for n in (1,2,3): make_template(n)

d=base('Django','Django Sample v1.1','Illustrative overlay · github.com/django/django')
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; run(p,'A pragmatic Python web framework — illustrative Four-Square overlay, not upstream branch practice.',7.8,True,NAVY)
lineage(d,'code','docs'); identity(d,'Represent a coherent Python web framework surface.','Keep code, public behavior and documentation aligned.','Illustrative mapping only; application-specific policy is outside the sample.')
table3(d,'Representative workflows',[('Request → response','routing + views + middleware','HTTP request follows a deterministic framework path.'),('Model ↔ database','ORM + migrations','Declared model state maps to database operations.'),('Template rendering','template engine','Context renders under documented escaping rules.'),('Forms / validation','forms + validators','Input is parsed and errors are surfaced consistently.'),('Testing','Django test utilities','Behavior can be exercised reproducibly.')])
lanes(d,'code','docs','Framework implementation, tests, migrations and runtime behavior.','Guides, tutorials, references, release notes and upgrade guidance.'); banner(d); page2(d,'code','docs')
d.save(OUT/'FSL-v1.1-Sample-Django-2P.docx')
print('built', OUT)
