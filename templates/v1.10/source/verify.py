#!/usr/bin/env python3
"""Verify the v1.10 source/document correspondence and optional package checksums.

This validates the delivered documents, not live project behavior or agent execution.
"""
from __future__ import annotations
import argparse
import hashlib
import re
from pathlib import Path, PurePosixPath
import fitz
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from build import AXIOM, VERSION_LINE, SPECS, NORM, blocks, plain


def require(condition: bool, message: str):
    if not condition: raise ValueError(message)


def compact(text: str) -> str:
    return re.sub(r'\s+', '', text.replace('\u00ad', ''))


def paragraphs(doc):
    for el in doc.element.body.iterchildren():
        if el.tag==qn('w:p'):
            p=Paragraph(el,doc)
            if p.text.strip(): yield p
        elif el.tag==qn('w:tbl'):
            t=Table(el,doc)
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        if p.text.strip(): yield p


def verify_documents(root: Path) -> list[str]:
    messages=[]
    for source,stem,identity,pages,grid in SPECS:
        doc=Document(root/(stem+'.docx'))
        expected=[plain(text) for kind,text in blocks(root/source)
                  if plain(text)!=AXIOM and not (grid and kind=='h3' and text=='Four squares')]
        ps=list(paragraphs(doc)); actual=[p.text for p in ps]
        require(actual==expected,f'{stem}: Markdown/DOCX body mismatch')
        require(len(doc.tables)==int(grid),f'{stem}: unexpected table')
        if grid:
            require(len(doc.tables[0].rows)==2 and len(doc.tables[0].columns)==2,
                    f'{stem}: grid must be four squares, not a run form')
        count=0
        for p in ps:
            for run in p.runs:
                if NORM.search(run.text):
                    require(run.bold and run.underline,f'{stem}: missing normative emphasis: {run.text}')
                    count+=1
        require(count>0 or 'APPENDIX' in stem, f'{stem}: no normative emphasis found')
        for s in doc.sections:
            foot=' '.join(p.text for p in s.footer.paragraphs)
            require(AXIOM in foot,f'{stem}: missing ownership footer')
            require(VERSION_LINE in foot,f'{stem}: missing template/draft footer')
            require(identity in ' '.join(p.text for p in s.header.paragraphs),
                    f'{stem}: missing separate project identity')
        pdf=fitz.open(root/(stem+'.pdf'))
        require(len(pdf)==pages,f'{stem}: expected {pages} pages, got {len(pdf)}')
        full=' '.join(page.get_text(sort=True) for page in pdf)
        for text in expected:
            # Multi-column ownership text interleaves in PDF sort order. Check each
            # grid paragraph using unsorted text extraction instead.
            if not compact(text) in compact(full):
                unsorted=' '.join(page.get_text() for page in pdf)
                require(compact(text) in compact(unsorted),f'{stem}: missing PDF source text: {text[:90]}')
        for index,page in enumerate(pdf,1):
            text=page.get_text()
            require(AXIOM in text and VERSION_LINE in text,f'{stem}/{index}: PDF identity/footer missing')
            require(identity in text,f'{stem}/{index}: project identity missing')
            require('\ufffd' not in text,f'{stem}/{index}: replacement glyph present')
            require(re.search(rf'\b{index}\s*/\s*{pages}\b',text),f'{stem}/{index}: page numbering mismatch')
            for b in page.get_text('blocks'):
                require(20<=b[0] and b[2]<=page.rect.width-20 and b[1]>=10 and b[3]<=page.rect.height-10,
                        f'{stem}/{index}: text outside safe page bounds')
        messages.append(f'PASS {stem}: {pages} page(s), exact body correspondence, {count} emphasized runs, correct grid count and per-page identity')
    return messages


def verify_policy(root: Path) -> list[str]:
    """Presence checks validate draft wording, never real enforcement."""
    method=plain((root/'FSL-v1.10-METHOD.md').read_text())
    anchors=(
      'single source of truth for its agreed business and technical intent',
      'Documents govern. Skills maintain. Evidence verifies.',
      'Intent approval, execution/spending permission and release acceptance are separate decisions',
      'an agent needs explicit delegation', 'FSL first inside SPEC',
      'one repository or several coordinated repositories',
      '[ SPEC | SYSTEM | AI | MEMORY ]',
      'Squares own artifacts; routes cross only the artifacts needed for the claimed outcome',
      'Route defines expected coverage; run records observed coverage',
      'No compulsory I-numbering', 'Historical and future iteration plans are optional',
      'New projects SHOULD begin with four E2E learning goals',
      'GitHub Actions are OFF by default',
      'general develop, commit or release request is not permission',
      'one start per rolling 60 minutes', 'explicit task total and expiry',
      'Two per hour needs justification and approval',
      'including failed, cancelled or uncertain attempts', 'an all-calls cap includes them',
      'Share the budget across the named repositories and agents',
      'not a claim that remote settings or hourly enforcement were configured',
      'Drift is a mismatch, not permission to rewrite the agreement',
      'Keep uncertain intent, authority, evidence or compatibility explicit',
      'next decision with its owner', 'Missing owners stay unassigned',
      'A scoped review may conclude no change is needed',
      'Analogies and homologies explain structural correspondences',
      'Appendix explains; errata corrects without changing intent; an adopted amendment changes intent',
      'A proposal is not approval', 'main1 is its original first commit',
      'CONFORMANT WITH EXCEPTION', 'MEMORY may preserve Project DNA; MEMORY cannot establish it',
      'each project establishes its own DNA in SPEC',
      'A sister cell MUST NOT silently become the parent',
      'only explicit Adoption of a SPEC revision creates a new DNA Generation',
      'Combining sister-cell variations creates a new candidate',
      'reuse evidence only where its recorded scope applies',
      'Copies, forks, deployments and handovers are different events',
      'Ancestry grants no current approval or execution permission',
      'Release edition. Project adoption remains separate.',
    )
    for token in anchors: require(token in method, f'Method missing: {token}')
    appendix=plain((root/'FSL-v1.10-APPENDIX-DNA-LINEAGE.md').read_text())
    for token in ('Explanatory only', 'Template adoption', 'Forking', 'Deployment', 'Handover',
                  'creates no new square', 'Projects may omit all biological vocabulary'):
        require(token in appendix, f'Appendix missing: {token}')
    skill=plain((root/'FSL-v1.10-SKILL.md').read_text())
    for token in ('released method', 'Shared ancestry is not inherited permission or proof',
                  'stop if consumption is unknown', 'Concurrency alone is not an hourly quota'):
        require(token in skill, f'Skill missing: {token}')
    example=plain((root/'FSL-v1.10-EXAMPLE-EVIDENCE-BRIEF.md').read_text())
    for token in ('invented teaching fixtures', 'No live fork, handover, agent, deployment or workflow was executed',
                  'cannot mark its own R-01 as PASS', 'marks the draft NON-CONFORMANT'):
        require(token in example, f'Example missing: {token}')
    return [f'PASS {len(anchors)} method anchors and appendix/skill/example boundaries (wording only)']


def verify_checksums(root: Path) -> int:
    manifest=(root/'MANIFEST.txt').read_text().splitlines()
    require(len(manifest)==len(set(manifest)),'Duplicate manifest entry')
    for name in manifest:
        p=PurePosixPath(name)
        require(not p.is_absolute() and '..' not in p.parts and '\\' not in name,'Unsafe manifest path')
    observed={p.relative_to(root).as_posix() for p in root.rglob('*')
              if p.is_file() and '__pycache__' not in p.parts and p.name not in ('MANIFEST.txt','SHA256SUMS.txt')}
    require(observed==set(manifest),'Manifest does not match payload')
    entries={}
    for line in (root/'SHA256SUMS.txt').read_text().splitlines():
        digest,name=line.split('  ',1)
        require(name not in entries,'Duplicate checksum entry')
        entries[name]=digest
    require(set(entries)==set(manifest)|{'MANIFEST.txt'},'Incomplete checksum coverage')
    for name,digest in entries.items():
        require(hashlib.sha256((root/name).read_bytes()).hexdigest()==digest,f'Checksum mismatch: {name}')
    return len(entries)

if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--root',type=Path,default=Path(__file__).resolve().parents[1])
    p.add_argument('--checksums',action='store_true')
    args=p.parse_args()
    for msg in verify_policy(args.root) + verify_documents(args.root): print(msg)
    if args.checksums: print(f'PASS {verify_checksums(args.root)} checksum entries')
