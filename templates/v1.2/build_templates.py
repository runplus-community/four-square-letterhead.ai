from pathlib import Path
import re, hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parent
NAVY='173D69'; BLUE='1F4E79'; LIGHT='F2F6FA'; PALE='E9F0F7'; GREEN='EAF4F1'; PURPLE='F3EEF8'; RED='FBECEC'; GOLD='F8F2E4'; BORDER='B7C2CF'; TEXT='222222'; MUTED='666666'
TOKENS=['CONFORMANT WITH EXCEPTION','NOT YET PROVEN','NON-CONFORMANT','SHOULD NOT','SHALL NOT','MUST NOT','CONFORMANT','EXCEPTION','WAIVER','SHOULD','SHALL','MUST','MAY']


def fmt(run,size=7,bold=False,color=TEXT,underline=False,italic=False):
    run.font.name='Aptos'; run.font.size=Pt(size); run.bold=bold; run.underline=underline; run.italic=italic
    run.font.color.rgb=RGBColor.from_string(color)


def norm(p,text,size=6.6,bold=False,color=TEXT):
    pat='('+'|'.join(re.escape(x) for x in sorted(TOKENS,key=len,reverse=True))+')'
    for part in re.split(pat,text):
        if not part: continue
        if part in TOKENS:
            r=p.add_run(part); fmt(r,size,True,NAVY,True)
        else:
            r=p.add_run(part); fmt(r,size,bold,color)


def shade(cell,fill):
    pr=cell._tc.get_or_add_tcPr(); sh=pr.find(qn('w:shd'))
    if sh is None: sh=OxmlElement('w:shd'); pr.append(sh)
    sh.set(qn('w:fill'),fill)


def margins(cell,top=40,start=45,bottom=40,end=45):
    pr=cell._tc.get_or_add_tcPr(); mar=pr.first_child_found_in('w:tcMar')
    if mar is None: mar=OxmlElement('w:tcMar'); pr.append(mar)
    for k,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        e=mar.find(qn('w:'+k))
        if e is None: e=OxmlElement('w:'+k); mar.append(e)
        e.set(qn('w:w'),str(v)); e.set(qn('w:type'),'dxa')


def borders(t,color=BORDER,sz='4'):
    pr=t._tbl.tblPr; bd=pr.first_child_found_in('w:tblBorders')
    if bd is None: bd=OxmlElement('w:tblBorders'); pr.append(bd)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=bd.find(qn('w:'+edge))
        if e is None: e=OxmlElement('w:'+edge); bd.append(e)
        e.set(qn('w:val'),'single'); e.set(qn('w:sz'),sz); e.set(qn('w:color'),color)


def widths(t,vals):
    for row in t.rows:
        for i,v in enumerate(vals): row.cells[i].width=Inches(v)


def heading(d,text):
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(1)
    r=p.add_run(text); fmt(r,8.2,True,NAVY); return p


def base_doc(project='[PROJECT / PRODUCT]',version='[PROJECT VERSION]'):
    d=Document(); s=d.sections[0]
    s.page_width=Inches(8.27); s.page_height=Inches(11.69); s.top_margin=Inches(.58); s.bottom_margin=Inches(.48); s.left_margin=Inches(.46); s.right_margin=Inches(.46)
    st=d.styles['Normal']; st.font.name='Aptos'; st.font.size=Pt(6.8); st.paragraph_format.space_after=Pt(.7)
    h=s.header; ht=h.add_table(rows=1,cols=2,width=Inches(7.25)); ht.alignment=WD_TABLE_ALIGNMENT.CENTER; ht.autofit=False; widths(ht,[5.1,2.15]); borders(ht,NAVY,'8')
    p=ht.cell(0,0).paragraphs[0]; r=p.add_run(project); fmt(r,12.3,True,NAVY); r=p.add_run('  /  FOUR-SQUARE CONSTITUTIONAL LETTERHEAD'); fmt(r,7.1,True,BLUE)
    p=ht.cell(0,1).paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT; r=p.add_run(version+'\n'); fmt(r,7,True,NAVY); r=p.add_run('Four-Square Template v1.2'); fmt(r,6.6,True,MUTED)
    f=s.footer; ft=f.add_table(rows=1,cols=3,width=Inches(7.25)); ft.alignment=WD_TABLE_ALIGNMENT.CENTER; ft.autofit=False; widths(ft,[2.55,2.15,2.55]); borders(ft,NAVY,'4')
    vals=['Four-Square Template v1.2\nSPEC + MEMORY fixed','MUST / SHOULD / MAY\nforce is explicit','main1 = immutable root\nmain = accepted integration']
    for i,x in enumerate(vals):
        p=ft.cell(0,i).paragraphs[0]; p.alignment=[WD_ALIGN_PARAGRAPH.LEFT,WD_ALIGN_PARAGRAPH.CENTER,WD_ALIGN_PARAGRAPH.RIGHT][i]; norm(p,x,5.5,False,MUTED)
    return d


def lineage(d,a='[lane-a]',b='[lane-b]',v='1.2'):
    heading(d,'FOUR-SQUARE RELEASE LINEAGE')
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('main1 = repository first commit / immutable root reference'); fmt(r,6.6,True,NAVY)
    t=d.add_table(rows=1,cols=4); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; widths(t,[1.8]*4); borders(t)
    vals=[f'SPEC\nspec-v{v}',f'{a.upper()}\n{a}-v{v}',f'{b.upper()}\n{b}-v{v}',f'MEMORY\nmemory-v{v}']
    fills=[LIGHT,PALE,GREEN,PURPLE]
    for i,x in enumerate(vals):
        c=t.cell(0,i); shade(c,fills[i]); margins(c); p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(x); fmt(r,6.6,True,NAVY if i in (0,3) else TEXT)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('SPEC and MEMORY are fixed. The two middle lanes are project-defined and SHOULD remain stable across versions.'); norm(p,r.text,6.1); r.text=''


def language(d):
    heading(d,'CONSTITUTIONAL LANGUAGE')
    t=d.add_table(rows=1,cols=5); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; widths(t,[1.45]*5); borders(t)
    items=[('MUST / MUST NOT','binding law',RED),('SHALL / SHALL NOT','constitutional alias',GOLD),('SHOULD / SHOULD NOT','strong guidance',PALE),('MAY','permitted option',GREEN),('EXCEPTION / WAIVER','explicit departure',PURPLE)]
    for i,(a,b,fill) in enumerate(items):
        c=t.cell(0,i); shade(c,fill); margins(c); p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; norm(p,a+'\n'+b,5.9)


def identity(d,mission='[MISSION]',principle='[CORE PRINCIPLE]',boundary='[BOUNDARY / OUT OF SCOPE]'):
    t=d.add_table(rows=1,cols=3); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; widths(t,[2.42]*3); borders(t)
    for i,(h,b) in enumerate([('MISSION',mission),('CORE PRINCIPLE',principle),('BOUNDARY',boundary)]):
        c=t.cell(0,i); shade(c,LIGHT if i!=1 else 'F8FAFC'); p=c.paragraphs[0]; r=p.add_run(h+'\n'); fmt(r,6.5,True,NAVY); norm(p,b,6.1)


def hierarchy(d):
    heading(d,'CONSTITUTIONAL HIERARCHY / FORMS')
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('CONSTITUTION  →  ARTICLES / LAWS  →  CONTRACTS  →  POLICIES  →  GUARDRAILS  →  RECOMMENDED PRACTICES  →  CHECKLISTS'); fmt(r,6.3,True,NAVY)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('The form label does not determine force; the normative keyword does.'); fmt(r,5.9,False,MUTED,False,True)


def laws(d,sample=False):
    heading(d,'ARTICLES / LAWS / CONTRACTS / GUARDRAILS')
    rows=(
      [('L1','LAW','MUST','Binding requirements MUST be explicit and observable.'),('L2','GUARDRAIL','MUST NOT','A binding requirement MUST NOT be silently weakened, bypassed or guessed away.'),('L3','POLICY','SHOULD','Material SHOULD deviations SHOULD have a recorded reason.'),('L4','CONTRACT','MUST','Missing evidence MUST result in NOT YET PROVEN, not assumed success.'),('L5','MEMORY','MUST NOT','MEMORY MUST NOT create new law unless promoted into SPEC.')]
      if not sample else
      [('D1','LAW','MUST','Declared public behavior MUST have implementation evidence and aligned developer guidance.'),('D2','GUARDRAIL','MUST NOT','DOCS MUST NOT knowingly describe behavior unsupported by accepted CODE.'),('D3','POLICY','SHOULD','Compatibility changes SHOULD record affected versions and migration guidance.'),('D4','CONTRACT','MUST','Release evidence MUST distinguish proven behavior from illustrative material.'),('D5','MEMORY','MUST NOT','MEMORY MUST NOT create new framework law unless promoted into SPEC.')]
    )
    t=d.add_table(rows=1+len(rows),cols=4); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; widths(t,[.65,1.05,1.05,4.5]); borders(t)
    for i,h in enumerate(['ID','FORM','FORCE','CLAUSE']): shade(t.cell(0,i),NAVY); p=t.cell(0,i).paragraphs[0]; r=p.add_run(h); fmt(r,5.8,True,'FFFFFF')
    for ri,row in enumerate(rows,1):
        for ci,x in enumerate(row):
            c=t.cell(ri,ci); shade(c,'FFFFFF' if ri%2 else 'F8FAFC'); p=c.paragraphs[0]; norm(p,x,5.65,bold=ci in (0,1))


def development(d):
    heading(d,'DEVELOPMENT GOVERNANCE — RECOMMENDED CHECKLIST')
    items=['Material work SHOULD use a checklist or equivalent visible plan.','Each binding clause SHOULD map to a verification item and observable evidence.','Agents MAY perform Discover → Plan → Execute → Observe → Verify → Record with defined responsibilities.','Agents MUST NOT reinterpret or waive constitutional law without explicit authority.','Proven libraries, binaries, scripts, documents, services, datasets, models, plugins and utilities SHOULD be evaluated before new core work.']
    for x in items:
        p=d.add_paragraph(); p.paragraph_format.left_indent=Inches(.12); r=p.add_run('□ '); fmt(r,6.1,True,NAVY); norm(p,x,6.0)


def conform(d):
    heading(d,'CONFORMANCE / EXCEPTION DECLARATION')
    t=d.add_table(rows=1,cols=4); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; widths(t,[1.8]*4); borders(t,NAVY,'5')
    for i,(x,fill) in enumerate([('CONFORMANT',GREEN),('CONFORMANT WITH EXCEPTION',PURPLE),('NON-CONFORMANT',RED),('NOT YET PROVEN',GOLD)]):
        c=t.cell(0,i); shade(c,fill); p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; norm(p,x,5.9,True)
    p=d.add_paragraph(); norm(p,'A MUST / MUST NOT deviation MUST be visible. An EXCEPTION / WAIVER MUST identify clause, scope, owner, rationale, evidence and review/expiry condition where applicable.',6.0)


def release(d):
    heading(d,'RELEASE LAW')
    p=d.add_paragraph(); norm(p,'Draft ≠ Tested ≠ Accepted ≠ Released. Claims MUST be scoped to evidence actually validated. A release MUST NOT hide unresolved binding violations. One version, four squares, all green.',6.2,True)


def page1(d,a='[lane-a]',b='[lane-b]',sample=False):
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('PROJECT CONSTITUTION / CONTRACT / LETTERHEAD'); fmt(r,8.2,True,NAVY)
    lineage(d,a,b); language(d)
    if sample: identity(d,'Provide a coherent framework for building database-backed web applications with mature integrated facilities.','CODE and DOCS SHOULD remain aligned; deviations SHOULD be explicit.','Application business logic and deployment policy remain with consuming projects.')
    else: identity(d)
    hierarchy(d); laws(d,sample); conform(d); release(d)


def page2(d,sample=False):
    heading(d,'PAGE 2 — GOVERNANCE / DEVELOPMENT / EXCEPTIONS')
    p=d.add_paragraph(); r=p.add_run('Use when project complexity requires more than the one-page constitutional identity.'); fmt(r,6.2,False,MUTED,False,True)
    laws(d,sample); development(d)
    heading(d,'EXCEPTION / WAIVER PROTOCOL')
    t=d.add_table(rows=2,cols=6); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; widths(t,[1.2]*6); borders(t)
    for i,h in enumerate(['CLAUSE','SCOPE','OWNER','RATIONALE','EVIDENCE','REVIEW / EXPIRY']): shade(t.cell(0,i),NAVY); p=t.cell(0,i).paragraphs[0]; r=p.add_run(h); fmt(r,5.5,True,'FFFFFF')
    for i in range(6): p=t.cell(1,i).paragraphs[0]; r=p.add_run('[REQUIRED]'); fmt(r,5.4,False,MUTED)
    conform(d); release(d)


def page3(d):
    heading(d,'PAGE 3 — AMENDMENT LAW / EVIDENCE LEDGER / CERTIFICATION')
    p=d.add_paragraph(); r=p.add_run('Use only for complex projects. Three pages is the maximum.'); fmt(r,6.2,False,MUTED,False,True)
    heading(d,'AMENDMENT LAW')
    for x in ['Binding constitutional changes MUST be made in SPEC and versioned. MEMORY MUST NOT amend the constitution.','A breaking constitutional change SHOULD identify migration impact.','Changing a project-selected middle lane token MUST be a reviewed migration, not an incidental rename.','Non-normative examples and mnemonics MAY evolve without changing constitutional force.']:
        p=d.add_paragraph(); norm(p,x,6.1)
    heading(d,'CONFORMANCE EVIDENCE LEDGER')
    t=d.add_table(rows=5,cols=5); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; widths(t,[.8,2.0,1.0,1.8,1.7]); borders(t)
    for i,h in enumerate(['ID','CLAUSE','FORCE','EVIDENCE','RESULT']): shade(t.cell(0,i),NAVY); p=t.cell(0,i).paragraphs[0]; r=p.add_run(h); fmt(r,5.7,True,'FFFFFF')
    for ri in range(1,5):
        for ci,x in enumerate([f'[{ri}]','[REQUIREMENT]','[MUST / SHOULD]','[TEST / REVIEW / HASH]','[STATE]']): p=t.cell(ri,ci).paragraphs[0]; norm(p,x,5.4)
    heading(d,'RELEASE CERTIFICATION')
    for x in ['Four lane versions MUST align to the declared release view.','No unresolved MUST / MUST NOT violation MAY be hidden.','Every active EXCEPTION MUST be attached to release evidence.','Every generated page MUST show Four-Square Template v1.2 separately from the project version.']:
        p=d.add_paragraph(); norm(p,x,6.1)


def make_template(n):
    d=base_doc(); page1(d)
    if n>=2: d.add_page_break(); page2(d)
    if n>=3: d.add_page_break(); page3(d)
    path=OUT/f'FSL-v1.2-Template-{n}P.docx'; d.save(path)


def make_django():
    d=base_doc('DJANGO / FOUR-SQUARE SAMPLE','Django Sample v1.2')
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Illustrative Four-Square constitutional overlay only — not upstream Django governance.'); fmt(r,6.3,False,MUTED,False,True)
    page1(d,'code','docs',True); d.add_page_break(); page2(d,True)
    path=OUT/'FSL-v1.2-Sample-Django-2P.docx'; d.save(path)


if __name__=='__main__':
    for n in (1,2,3): make_template(n)
    make_django()
    print('built FSL v1.2 DOCX templates and Django sample')
