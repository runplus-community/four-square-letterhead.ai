#!/usr/bin/env python3
"""Generate the FSL v1.7 editable letterheads; render these DOCX files to PDF.
Standalone: Python 3.10+ and python-docx 1.2.0. No prior release or network needed.
"""
from __future__ import annotations
import argparse
import re
from datetime import datetime, timezone
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

VERSION = '1.7'
AXIOM = 'FSL owns the method. The project owns the work.'
FONT = 'Liberation Sans'
INK, NAVY, MUTED, LINE, PALE = '172B3A', '183B54', '526474', 'C9D3DB', 'F2F5F7'
WIDTH = 510
TERMS = ('CONFORMANT WITH EXCEPTION', 'NOT YET PROVEN', 'NON-CONFORMANT',
         'MUST NOT', 'SHOULD NOT', 'SHALL NOT', 'MUST', 'SHOULD', 'SHALL', 'MAY',
         'CONFORMANT', 'EXCEPTION', 'WAIVER')
NORM = re.compile(r'(?<![A-Za-z0-9_-])(' + '|'.join(map(re.escape, TERMS)) + r')(?![A-Za-z0-9_-])')


def run(p, text, size=8.8, bold=False, color=INK, underline=False):
    r = p.add_run(text)
    r.font.name = FONT; r.font.size = Pt(size * 1.10)
    r.bold = bold; r.underline = underline
    r.font.color.rgb = RGBColor.from_string(color)
    return r


def rich(p, text, size=8.8, bold=False, color=INK):
    for part in NORM.split(text):
        if part:
            run(p, part, size, bold or part in TERMS, color, part in TERMS)
    return p


def para(parent, text='', size=8.8, before=0, after=3, bold=False, color=INK, keep=False):
    p = parent.add_paragraph()
    f = p.paragraph_format
    f.space_before = Pt(before); f.space_after = Pt(after); f.line_spacing = 1.04
    f.keep_together = True; f.keep_with_next = keep
    return rich(p, text, size, bold, color)


def label(d, text, before=7):
    p = para(d, text, 9.3, before, 4, True, NAVY, True)
    p.style = d.styles['Heading 2']
    return p


def shade(c, fill):
    el = OxmlElement('w:shd'); el.set(qn('w:fill'), fill)
    c._tc.get_or_add_tcPr().append(el)


def table(d, widths, rows, header=None, size=8.5, boxes=False):
    t = d.add_table(rows=0, cols=len(widths)); t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    w = t._tbl.tblPr.find(qn('w:tblW'))
    w.set(qn('w:w'), str(int(sum(widths) * 20))); w.set(qn('w:type'), 'dxa')
    for c, width in zip(t.columns, widths): c.width = Pt(width)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top','bottom','left','right','insideH','insideV'):
        el=OxmlElement('w:'+edge)
        el.set(qn('w:val'), 'single' if boxes or edge in ('top','bottom','insideH') else 'nil')
        el.set(qn('w:sz'),'5'); el.set(qn('w:color'),LINE); borders.append(el)
    t._tbl.tblPr.append(borders)
    for n, values in enumerate(([header] if header else []) + list(rows)):
        row=t.add_row(); pr=row._tr.get_or_add_trPr(); pr.append(OxmlElement('w:cantSplit'))
        if header and n==0: pr.append(OxmlElement('w:tblHeader'))
        for c, width, val in zip(row.cells, widths, values):
            c.width=Pt(width); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            mar=OxmlElement('w:tcMar')
            for edge, value in (('top',70),('bottom',70),('start',100),('end',100)):
                el=OxmlElement('w:'+edge); el.set(qn('w:w'),str(value)); el.set(qn('w:type'),'dxa'); mar.append(el)
            c._tc.get_or_add_tcPr().append(mar)
            p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.04
            rich(p,val,7.7 if header and n==0 else size,bool(header and n==0),NAVY if header and n==0 else INK)
            if header and n==0: shade(c,PALE)
    return t


def field(p, name):
    el=OxmlElement('w:fldSimple'); el.set(qn('w:instr'),name)
    r=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
    sz=OxmlElement('w:sz'); sz.set(qn('w:val'),'15'); pr.append(sz); r.append(pr)
    tx=OxmlElement('w:t'); tx.text='1'; r.append(tx); el.append(r); p._p.append(el)


def source_link(d, text, url):
    p=para(d,'',8.3,after=4)
    el=OxmlElement('w:hyperlink'); el.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
    size=OxmlElement('w:sz'); size.set(qn('w:val'),'16'); pr.append(size)
    color=OxmlElement('w:color'); color.set(qn('w:val'),NAVY); pr.append(color); r.append(pr)
    tx=OxmlElement('w:t'); tx.text=text; r.append(tx); el.append(r); p._p.append(el)
    run(p,'  '+url.replace('https://',''),7.4,color=MUTED)


def new_document(sample=False):
    d=Document(); s=d.sections[0]
    s.page_width=Pt(595.28); s.page_height=Pt(841.89)
    s.left_margin=s.right_margin=Pt((595.28-WIDTH)/2)
    s.top_margin=Pt(88); s.bottom_margin=Pt(56)
    s.header_distance=Pt(24); s.footer_distance=Pt(21)
    normal=d.styles['Normal']; normal.font.name=FONT; normal.font.size=Pt(8.8)
    normal.paragraph_format.space_after=Pt(3); normal.paragraph_format.line_spacing=1.04
    for name in ('Heading 1','Heading 2'):
        d.styles[name].font.name=FONT; d.styles[name].font.color.rgb=RGBColor.from_string(NAVY)
        d.styles[name].paragraph_format.keep_with_next=True
    d.core_properties.title=f'FSL v{VERSION} - '+('Django adoption illustration' if sample else 'Project letterhead template')
    d.core_properties.author='FSL'; d.core_properties.last_modified_by='FSL'
    d.core_properties.created=d.core_properties.modified=datetime(2026,9,9,tzinfo=timezone.utc)
    p=s.header.paragraphs[0]; p.paragraph_format.space_after=Pt(4)
    p.paragraph_format.tab_stops.add_tab_stop(Pt(WIDTH),WD_TAB_ALIGNMENT.RIGHT)
    run(p,'FSL / FOUR-SQUARE LETTERHEAD',8,True,NAVY)
    run(p,'\t'+('Sample edition v1.7' if sample else 'Project / contract version [vN]'),8,True,MUTED)
    para(s.header,'DJANGO' if sample else '[PROJECT / PRODUCT]',21,after=2,bold=True,color=NAVY)
    p=para(s.header,'ILLUSTRATIVE ADOPTION / reference: django/django' if sample else 'Project repository: [owner/repository]  /  Owner: [responsible person or team]',8,after=4,color=MUTED)
    borders=OxmlElement('w:pBdr'); b=OxmlElement('w:bottom')
    for key,val in (('val','single'),('sz','12'),('color',NAVY),('space','6')): b.set(qn('w:'+key),val)
    borders.append(b); p._p.get_or_add_pPr().append(borders)
    p=s.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(4)
    run(p,AXIOM,8.3,True,NAVY)
    p=para(s.footer,'',7.5,after=0)
    p.paragraph_format.tab_stops.add_tab_stop(Pt(WIDTH),WD_TAB_ALIGNMENT.RIGHT)
    run(p,f'Four-Square Template v{VERSION}',7.5,color=MUTED)
    run(p,'\tPage ',7.5,color=MUTED); field(p,'PAGE'); run(p,' / ',7.5,color=MUTED); field(p,'NUMPAGES')
    return d


def squares(d, sample=False):
    data=[('SPEC','spec','FSL / Scope / Contracts / Acceptance / Routes','Optional: Appendix / Errata / Amendments'),
          ('CODE' if sample else 'SYSTEM','code' if sample else 'system','Product / Runtime / Data / API / UI / Build','Gate: deliverable and product checks'),
          ('DOCS' if sample else 'AI','docs' if sample else 'ai','Developer guidance / Examples' if sample else 'Instructions / Agents / Skills / Prompts / Evals','Gate: aligned guidance' if sample else 'Gate: usable capability; declare inactivity'),
          ('MEMORY','memory','Decisions / Rationale / Context / Evidence links','Gate: reliable recall; no new law')]
    t=table(d,[255,255],[['',''],['','']],boxes=True)
    for i,(name,branch,items,gate) in enumerate(data):
        c=t.cell(i//2,i%2); shade(c,PALE if i in (0,3) else 'FFFFFF')
        p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(4)
        run(p,f'{i+1:02}  {name}',10,True,NAVY); run(p,f'  /  {branch}-vN',7.8,color=MUTED)
        para(c,items,8.3,after=3); para(c,gate,8,after=1,color=MUTED)
    para(d,'SPEC + MEMORY fixed; SYSTEM + AI recommended. '+('CODE / DOCS retained as this sample specialization.' if sample else 'Replace middle lanes only for a durable concern; ownership follows purpose, not author.'),7.7,before=3,after=1,color=MUTED)


def page_one(d, sample=False):
    para(d,'PURPOSE  '+('Illustrate adoption of an existing database-backed web application. [D1]' if sample else '[Who this serves and the observable outcome it promises.]'),9,after=3)
    para(d,'BOUNDARY  '+('Teaching only; not upstream Django policy, adoption, or tested behavior.' if sample else '[In scope / out of scope; the boundaries of the contract and proof.]'),8.8,after=2)
    label(d,'FOUR SQUARES',5); squares(d,sample)
    label(d,'BINDING CONTRACT',6)
    para(d,'MUST / MUST NOT = binding  /  SHOULD / SHOULD NOT = guidance  /  MAY = permission',7.5,after=4)
    rows=[('LAW-DJANGO-001','Public behavior MUST have scoped code evidence and aligned guidance.','Code + guidance'),('GUARD-DJANGO-002','DOCS MUST NOT claim behavior unsupported by accepted CODE.','Review'),('CONTRACT-DJANGO-004','Evidence MUST distinguish tested behavior from illustration.','Scoped evidence')] if sample else [('LAW-[SCOPE]-001','[Subject] MUST [observable requirement].','[Check / evidence]'),('CONTRACT-[SCOPE]-002','[Interface] MUST [observable promise].','[Acceptance check]'),('GUARD-[SCOPE]-003','[Operation] MUST NOT [forbidden effect].','[Boundary test]')]
    table(d,[121,282,107],rows,['CLAUSE ID','REQUIREMENT / PROMISE','VERIFICATION'],8.3)
    label(d,'E2E ROUTES / VISIBLE TOUCHPOINTS',6)
    path='request -> URL -> view -> model -> DB -> template -> HTML response' if sample else '[input] -> [touchpoint 1] -> [touchpoint 2] -> [observable output]'
    para(d,'R-01  '+path,8.5,bold=True,after=3)
    para(d,'Expected: '+('known article title in the returned HTML; browser/deployment excluded.' if sample else '[asserted outcome + boundary]. Route baseline: [SPEC/version or exact reference].'),7.9,after=2)
    para(d,'Every run shows every declared point: PASS / FAIL / NOT RUN + evidence. Listing points or passing a component test is not E2E proof.',7.8,after=1,color=MUTED)
    label(d,'ADOPTION / DEVELOPMENT STRATEGY',6)
    if sample:
        para(d,'Existing: study -> reviewed adoption baseline -> verify R-01. Iterations optional; do not reconstruct history. Baseline and runtime evidence not supplied.',8.3,after=1)
    else:
        para(d,'New: I-001 -> tiny R-01 -> verify -> grow. Existing: study -> adoption baseline -> verify R-01; iterations optional. Parallel / E2M / M2E MAY help.',8.3,after=1)
    label(d,'ESSENTIAL TECHNOLOGIES / REFERENCES',6)
    refs=[('Django facilities','SHOULD REUSE within this illustration','[D1] Overview'),('Django tests','REFERENCE for planned checks','[D2] Testing')] if sample else [('[Technology / library]','[MUST / SHOULD / MAY + use]','[Canonical source / clause]'),('[Proven capability]','[REUSE / INHERIT / REFERENCE]','[Source; pin if material]')]
    table(d,[120,239,151],refs,['ITEM','INTENT / USE','SOURCE'],8.1)
    label(d,'LIFECYCLE / RELEASE CHECK',6)
    para(d,'Draft -> Prove -> Release -> Maintain as needed -> Next release. Appendix explains; errata corrects; adopted amendment changes intent. Released SPEC stays frozen.',7.8,after=3)
    a,b=('CODE','DOCS') if sample else ('SYSTEM','AI')
    para(d,f'[ ] SPEC stable   [ ] {a} verified   [ ] {b} aligned / scoped   [ ] MEMORY understandable',7.8,after=3)
    para(d,'Conformance: NOT YET PROVEN. '+('Example only; no runtime run or release snapshot asserted.' if sample else 'Effective SPEC / adopted changes / four-commit snapshot: [project-local reference].'),7.7,after=2)
    para(d,'main1 = original root. New lanes start there; existing adoption may start at a reviewed commit. Matching lane lineage continues; main = accepted integration.',7.4,after=0,color=MUTED)


def next_page(d,title,sub):
    d.add_page_break(); p=para(d,title,16,after=6,bold=True,color=NAVY); p.style=d.styles['Heading 1']
    para(d,sub,8.8,after=8,color=MUTED)


def page_two(d):
    next_page(d,'ROUTES, RUNS & ADOPTION','Use this page when the primary path, proof boundary or adoption context needs room. Omit unused detail.')
    label(d,'R-01 / ROUTE DEFINITION',3)
    para(d,'Route baseline: [SPEC/version + exact reference]. Clauses: [IDs]. Boundary: [input -> output]. Expected result: [observable assertion].',9,after=5)
    para(d,'[input] -> [touchpoint 1] -> [touchpoint 2] -> [output]',9,bold=True,after=5)
    para(d,'A route is stable within its baseline; a run is an execution, not an iteration. Add R-02 only for another critical path. Do not shrink a route to match a passing subset.',8.7,after=3)
    label(d,'REUSABLE RUN SHEET / SHOW BEFORE EXECUTION')
    para(d,'Run: [unique ID/date]. Implementation: [commit/artifact]. Environment: [host/versions]. SPEC + adopted changes: [exact references]. Route: R-01 at [baseline].',8.6,after=5)
    table(d,[116,178,79,137],[('[input]','[Observed input/event]','NOT RUN','[Evidence]'),('[touchpoint 1]','[Expected participation]','NOT RUN','[Evidence]'),('[touchpoint 2]','[Expected participation]','NOT RUN','[Evidence]'),('[output]','[Expected visible result]','NOT RUN','[Evidence]')],['REQUIRED TOUCHPOINT','OBSERVATION / ASSERTION','RESULT','EVIDENCE'],8.5)
    para(d,'Outcome: [expected / observed]. Route result: NOT YET PROVEN. Every point needs evidence from the connected workflow; any unreached point remains NOT RUN.',8.7,before=5,after=4)
    para(d,'A negative route may PASS when its expected rejection is verified. Fixtures, stubs, E2M and M2E prove only the declared scope. Do not prefill success or reuse a prior run result automatically.',8.6,after=3)
    label(d,'EXISTING PROJECT / ADOPT THE PRESENT')
    para(d,'Study code, documentation and runtime as needed. Record the current repository/commit, scope and known gaps. Separate OBSERVED / AS-IS from ADOPTED / INTENDED behavior. A bug does not become law because it exists.',9,after=5)
    para(d,'Iterations are optional for existing projects, including future iteration plans. Preserve history and justified lane names. Prove representative routes only for the claims made; a partial adoption baseline is not universal certification.',9,after=3)
    label(d,'MIDPOINTS, OWNERSHIP & REUSE')
    para(d,'M = [named contract/interface]. Parallel paths must agree there; verify their shared state or integration. A runtime route crosses AI only when AI actually participates. Unchanged touchpoints can be verified without edits.',8.8,after=5)
    table(d,[98,412],[('REUSE','Use a capability without automatically adopting its source law.'),('INHERIT','Explicitly adopt a named requirement/version/scope into project SPEC.'),('REFERENCE','Consult knowledge/evidence without adopting its normative authority.')],size=8.7)
    para(d,'One authoritative owner per artifact. Agent-written product code is SYSTEM; AI instructions are AI; adopted promises are SPEC; rationale and handover are MEMORY.',8.6,before=5,after=2,color=MUTED)
    label(d,'SUPPORTING SPEC / ONLY WHEN USEFUL')
    para(d,'Appendix: [optional study/diagram/example]. Errata: [optional correction]. Amendments: [optional adopted changes]. No mandatory empty documents, YAML, route database or iteration ledger.',8.6,after=0)


def page_three(d):
    next_page(d,'BASELINE, LIFECYCLE & RELEASE','A reviewed combination of four squares, not four unrelated green checks. Keep exact evidence in project-local files.')
    label(d,'EXACT FOUR-SQUARE SNAPSHOT',3)
    para(d,'Project / release: [identity]. Repository: [owner/repository]. FSL template: v1.7.',8.8,after=5)
    table(d,[77,127,176,130],[('SPEC','spec-vN','[Full reviewed commit]','[Gate evidence]'),('SYSTEM','system-vN','[Full reviewed commit]','[Gate evidence]'),('AI','ai-vN','[Full reviewed commit]','[Gate / inactive scope]'),('MEMORY','memory-vN','[Full reviewed commit]','[Gate evidence]')],['SQUARE','PROJECT LANE','REVIEWED COMMIT','EVIDENCE'],8.5)
    para(d,'Integration: [exact commit]. Route evidence: [R-01/run reference]. Publication: [actual branch/tag/release state]. Unchanged lanes need review, not ceremonial edits. Git remains authoritative.',8.6,before=5,after=3)
    label(d,'EFFECTIVE CONTRACT / FROZEN BASELINE')
    para(d,'Released SPEC: [version + immutable reference]. Applicable adopted changes: [IDs + immutable references, or NONE]. Adoption decision / effective scope: [reference]. A committed proposal is not automatically adopted.',8.9,after=5)
    table(d,[100,410],[('APPENDIX','Supporting explanation before/during development or after release; no normative change.'),('ERRATA','Correct an error without changing normative intent.'),('AMENDMENT','Change normative intent explicitly: base version, clauses, exact delta, authorized adoption, effective scope/date and verification impact.')],size=8.7)
    para(d,'Before first release, reviewed draft SPEC may change directly. After release, keep the baseline separate and frozen. Resolve overlapping amendments before claiming conformance; record incorporated changes at the next release.',8.7,before=5,after=3)
    label(d,'CONFORMANCE / NOT A TEST STATUS')
    table(d,[164,346],[('CONFORMANT','Evidence supports all applicable binding requirements in scope.'),('CONFORMANT WITH EXCEPTION','An authorized scoped exception is visible; it is not a test pass.'),('NON-CONFORMANT','A binding requirement is known to be violated.'),('NOT YET PROVEN','Evidence or a necessary interpretation is unresolved.')],size=8.7)
    para(d,'Declared state: [NOT YET PROVEN]. Reviewer/date: [authority/date]. Exceptions: [NONE / IDs, scope, reason, evidence, review/expiry].',8.7,before=5,after=3)
    label(d,'HANDOVER / NEXT CLEAN BASELINE')
    para(d,'Adopted intent is not implemented behavior; implementation is not verified behavior. Reassess affected routes after a change. Historical evidence retains its original scope. MEMORY may remember law; MEMORY cannot make law.',9,after=5)
    para(d,'Release is immutable. Clarify with appendix, correct with errata, change with amendment.',8.8,bold=True,after=0)


def sample_page_two(d):
    next_page(d,'DJANGO / EXISTING-SYSTEM ILLUSTRATION','Sample edition v1.7 / reference baseline: Django 5.2 documentation. No actual upstream adoption or execution is asserted.')
    label(d,'R-01 / ONE ARTICLE TO HTML',3)
    para(d,'Illustrative route: a request retrieves one known stored article and returns its title in HTML. The request-to-response boundary excludes a browser, JavaScript, hosting and production deployment. [D1]',9,after=5)
    para(d,'request -> URL -> view -> model -> DB -> template -> HTML response',9,bold=True,after=4)
    label(d,'EVERY RUN SEES THE TOUCHPOINTS')
    para(d,'Run: NOT EXECUTED. Implementation/environment: not supplied. SPEC/route baseline: this illustrative sample only. No runtime evidence is attached.',8.7,after=5)
    table(d,[114,190,79,127],[('Request','Known article request','NOT RUN','No runtime evidence'),('URL','Resolve expected view','NOT RUN','No runtime evidence'),('View','Handle known arguments','NOT RUN','No runtime evidence'),('Model','Select expected record','NOT RUN','No runtime evidence'),('Database','Read known stored title','NOT RUN','No runtime evidence'),('Template','Render the retrieved title','NOT RUN','No runtime evidence'),('HTML response','Assert title in output','NOT RUN','No runtime evidence')],['TOUCHPOINT','PLANNED OBSERVATION','RESULT','EVIDENCE'],8.4)
    para(d,'Expected outcome: known title in returned HTML. Observed outcome: not supplied. Route conformance: NOT YET PROVEN. Documentation and page-layout checks are not runtime proof.',8.8,before=5,after=3)
    label(d,'PROSPECTIVE ADOPTION / NO INVENTED ITERATIONS')
    para(d,'Study the existing repository and deployed behavior; choose an exact adoption baseline; distinguish observed behavior from intended promises. Retain the CODE / DOCS specialization if justified. No historical or future iteration plan is required.',8.9,after=5)
    para(d,'AI is not an artificial runtime hop. An agent helping with this study would not prove the request path. A database fixture supports only its disclosed test scope, not a production integration.',8.7,after=3)
    label(d,'OPTIONAL SPEC LIFECYCLE')
    para(d,'Appendix may hold the current-state map. Errata can correct a recorded reference. A new post-release promise needs an explicitly adopted amendment tied to its baseline and clauses; the prior release stays frozen.',8.9,after=3)
    label(d,'CANONICAL REFERENCES / NOT ADOPTED UPSTREAM LAW')
    source_link(d,'[D1] Django 5.2 overview','https://docs.djangoproject.com/en/5.2/intro/overview/')
    source_link(d,'[D2] Django 5.2 testing','https://docs.djangoproject.com/en/5.2/topics/testing/overview/')
    para(d,'References checked 9 September 2026. Proposed route, ownership and sample clauses are FSL-authored teaching material, not Django governance. [D2] informs planned checks only.',8.2,before=4,after=0,color=MUTED)


def build(output: Path):
    output.mkdir(parents=True,exist_ok=True)
    for pages in (1,2,3):
        d=new_document(); page_one(d)
        if pages>=2: page_two(d)
        if pages>=3: page_three(d)
        p=output/f'FSL-v{VERSION}-Template-{pages}P.docx'; d.save(p); print(p.name)
    d=new_document(True); page_one(d,True); sample_page_two(d)
    p=output/f'FSL-v{VERSION}-Sample-Django-2P.docx'; d.save(p); print(p.name)


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--output',type=Path,default=Path.cwd())
    args=p.parse_args(); build(args.output)

if __name__=='__main__': main()
