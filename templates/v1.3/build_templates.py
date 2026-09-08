from pathlib import Path
from docx import Document
import re, shutil

SRC=Path('templates/v1.2')
OUT=Path('templates/v1.3')
OUT.mkdir(parents=True, exist_ok=True)

NORM_TERMS=[
    'CONFORMANT WITH EXCEPTION','NOT YET PROVEN','NON-CONFORMANT','SHOULD NOT','SHALL NOT','MUST NOT',
    'CONFORMANT','EXCEPTION','WAIVER','SHOULD','SHALL','MUST','MAY'
]
pat=re.compile(r'('+'|'.join(re.escape(x) for x in NORM_TERMS)+r')')

def set_para(p,text, base_bold=None):
    style=p.style
    p.text=''
    if style: p.style=style
    for part in pat.split(text):
        if not part: continue
        r=p.add_run(part)
        if base_bold is not None: r.bold=base_bold
        if part in NORM_TERMS:
            r.bold=True; r.underline=True
    return p

def set_cell(cell,text):
    cell.text=''
    lines=text.split('\n')
    set_para(cell.paragraphs[0],lines[0])
    for line in lines[1:]: set_para(cell.add_paragraph(),line)

def replace_version_everywhere(doc, sample=False):
    for p in doc.paragraphs:
        if 'v1.2' in p.text: set_para(p,p.text.replace('v1.2','v1.3'))
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                if 'v1.2' in c.text: set_cell(c,c.text.replace('v1.2','v1.3'))
    for s in doc.sections:
        for hf in (s.header,s.footer):
            for p in hf.paragraphs:
                for r in p.runs: r.text=r.text.replace('v1.2','v1.3')
            for t in hf.tables:
                for row in t.rows:
                    for c in row.cells:
                        for p in c.paragraphs:
                            for r in p.runs: r.text=r.text.replace('v1.2','v1.3')
    if sample:
        for s in doc.sections:
            for t in s.header.tables:
                for row in t.rows:
                    for c in row.cells:
                        for p in c.paragraphs:
                            for r in p.runs: r.text=r.text.replace('Django Sample v1.2','Django Sample v1.3')

def update_common(doc):
    if len(doc.paragraphs)>2:
        set_para(doc.paragraphs[2],'Fixed lanes: SPEC + MEMORY. Middle lanes are project-defined and version-stable. Framework/template version and governed project version MUST remain separate identities.')
    if len(doc.tables)>=4:
        t=doc.tables[3]
        set_cell(t.cell(0,1),'RELATION / EXECUTION SOURCE')
        rels=['REUSE: [PROVEN CAPABILITY]','INHERIT: [REQUIREMENT / CONTRACT]','REFERENCE: [EVIDENCE / KNOWLEDGE]','REUSE: [ASSET / SERVICE]','REFERENCE: [GUIDANCE / DATA]']
        for i,rel in enumerate(rels,1):
            if i<len(t.rows): set_cell(t.cell(i,1),rel)
    if len(doc.tables)>=6:
        set_cell(doc.tables[5].cell(0,3),'MEMORY\nmemory-vN\nHuman+AI recall, glossary, mnemonics, examples and maps. MEMORY may remember law; MEMORY cannot make law.')
    if len(doc.tables)>=7:
        t=doc.tables[6]
        set_cell(t.cell(0,0),'TRACEABILITY\nMaterial binding clauses MUST carry stable clause IDs and SHOULD map to scoped evidence.')
        set_cell(t.cell(0,1),'SNAPSHOT\nAccepted Four-Square releases MUST record the exact four lane commits. The snapshot MUST NOT replace Git.')
        set_cell(t.cell(0,2),'AGENTS\nAgents MAY assist, but MUST NOT invent law or silently choose a stronger/weaker meaning when normative intent is ambiguous.')
    if len(doc.tables)>=8:
        set_cell(doc.tables[7].cell(1,0),'RELEASE LAW - ONE VERSION, FOUR SQUARES, ALL GREEN\nA release MUST NOT claim conformance while a binding violation is unresolved. The four reviewed commit IDs MUST be recorded in the release snapshot.')

# 1P
p1=OUT/'FSL-v1.3-Template-1P.docx'
shutil.copy2(SRC/'FSL-v1.2-Template-1P.docx',p1)
d=Document(p1); replace_version_everywhere(d); update_common(d)
set_para(d.paragraphs[0],'[SHORT PRODUCT IDENTITY / CONSTITUTIONAL TAGLINE]')
set_para(d.paragraphs[4],'CANONICAL WORKFLOWS / CONTRACTS / RELATIONS')
set_para(d.paragraphs[7],'TRACEABILITY / RELEASE GOVERNANCE')
set_cell(d.tables[1].cell(0,4),'EXCEPTION / WAIVER\nExplicit scoped departure. Clause IDs and evidence keep deviations traceable.')
set_cell(d.tables[4].cell(0,5),'PRACTICE\nRecommended method; normally SHOULD / MAY. Supporting artifacts carry detail beyond the 1-3 page letterhead.')
for s in d.sections:
    if s.footer.tables: set_cell(s.footer.tables[0].cell(0,1),'Traceability\nCLAUSE ID -> EVIDENCE')
d.save(p1)

# 2P
p2=OUT/'FSL-v1.3-Template-2P.docx'
shutil.copy2(SRC/'FSL-v1.2-Template-2P.docx',p2)
d=Document(p2); replace_version_everywhere(d); update_common(d)
set_para(d.paragraphs[4],'CANONICAL WORKFLOWS / CONTRACTS / RELATIONS')
set_para(d.paragraphs[7],'TRACEABILITY / RELEASE GOVERNANCE')
set_para(d.paragraphs[9],'PAGE 2 - TRACEABILITY / VERSION / PROOF')
set_para(d.paragraphs[10],'Use Page 2 when clause identity, proof scope, version identity, release snapshot or authority boundaries need room. FSL remains constitutional; detailed logs and schemas stay in supporting artifacts.')
set_para(d.paragraphs[11],'ARTICLE I - CLAUSE IDENTITY / VERSION IDENTITY')
set_para(d.paragraphs[12],'ARTICLE II - PROOF / RELEASE SNAPSHOT')
set_para(d.paragraphs[13],'REUSE / INHERIT / REFERENCE')
set_para(d.paragraphs[14],'AI / MEMORY AUTHORITY')
bullets=[
    '• Material binding clauses MUST carry stable IDs such as LAW-SPEC-001 or CONTRACT-AI-002; a published ID MUST NOT be reused for a different meaning.',
    '• Four-Square Template v1.3 identifies the FSL governance language. [PROJECT VERSION] identifies the governed system. They MUST NOT be conflated.',
    '• An accepted Four-Square release MUST record the exact four lane commit IDs in a lightweight snapshot; Git remains the authority.',
    '• Proof SHOULD be recorded as CLAIM / SCOPE / ENVIRONMENT / EVIDENCE / RESULT. A result MUST NOT be generalized beyond the scope actually validated.',
    '• MEMORY may remember law; MEMORY cannot make law. An agent MUST surface ambiguous normative intent for review and MUST NOT silently choose the stronger or weaker interpretation.'
]
for idx,text in enumerate(bullets,15): set_para(d.paragraphs[idx],text)
T=d.tables[8]
rows=[
    ('CLAUSE ID','Stable identity','Material binding clauses MUST use a stable FORM-SCOPE-NNN identifier.'),
    ('VERSION IDENTITY','Separate governance/system versions','FSL template version and project version MUST be shown separately and MUST NOT be conflated.'),
    ('RELEASE SNAPSHOT','Four commit record','Accepted Four-Square releases MUST record exact SPEC / lane A / lane B / MEMORY commit IDs.'),
    ('PROOF TUPLE','Scoped evidence record','Evidence SHOULD state CLAIM / SCOPE / ENVIRONMENT / EVIDENCE / RESULT.'),
    ('MEMORY AXIOM','Authority boundary','MEMORY may remember law; MEMORY cannot make law.'),
    ('AI AMBIGUITY','Review boundary','An agent MUST NOT silently resolve ambiguous normative intent to a stronger or weaker rule.')
]
for r,row in enumerate([('FORM','PURPOSE','RULE')]+rows):
    for c,val in enumerate(row): set_cell(T.cell(r,c),val)
T=d.tables[9]
relrows=[
    ('REUSE','Use an existing proven capability. Reuse does not automatically import the source project\'s law or assumptions.'),
    ('INHERIT','Explicitly adopt an existing requirement or contract. Inherited law MUST be identified as adopted.'),
    ('REFERENCE','Use knowledge, evidence or prior decisions without adopting their normative authority.'),
    ('AMBIGUITY','When intent is unclear, an agent MUST surface it for review; it MUST NOT silently strengthen or weaken the rule.')
]
for r,(a,b) in enumerate(relrows): set_cell(T.cell(r,0),a); set_cell(T.cell(r,1),b)
T=d.tables[10]
set_cell(T.cell(0,0),'PROOF\nCLAIM / SCOPE / ENVIRONMENT / EVIDENCE / RESULT SHOULD accompany material conformance claims.')
set_cell(T.cell(0,1),'SNAPSHOT\nSPEC / lane A / lane B / MEMORY commit IDs MUST identify the reviewed release view.')
set_cell(T.cell(0,2),'VERSION\nTemplate version = FSL language. Project version = governed system. They MUST NOT be conflated.')
for s in d.sections:
    if s.footer.tables: set_cell(s.footer.tables[0].cell(0,1),'Traceability\nCLAUSE ID -> PROOF -> SNAPSHOT')
d.save(p2)

# 3P
p3=OUT/'FSL-v1.3-Template-3P.docx'
shutil.copy2(SRC/'FSL-v1.2-Template-3P.docx',p3)
d=Document(p3); replace_version_everywhere(d); update_common(d)
set_para(d.paragraphs[4],'CANONICAL WORKFLOWS / CONTRACTS / RELATIONS')
set_para(d.paragraphs[7],'TRACEABILITY / RELEASE GOVERNANCE')
set_para(d.paragraphs[9],'PAGE 2 - TRACEABILITY / VERSION / PROOF')
set_para(d.paragraphs[10],'Use Page 2 when clause identity, proof scope, version identity, release snapshot or authority boundaries need room. FSL remains constitutional; detailed logs and schemas stay in supporting artifacts.')
set_para(d.paragraphs[11],'ARTICLE I - CLAUSE IDENTITY / VERSION IDENTITY')
set_para(d.paragraphs[12],'ARTICLE II - PROOF / RELEASE SNAPSHOT')
set_para(d.paragraphs[13],'REUSE / INHERIT / REFERENCE')
set_para(d.paragraphs[14],'AI / MEMORY AUTHORITY')
for idx,text in enumerate(bullets,15): set_para(d.paragraphs[idx],text)
T=d.tables[8]
for r,row in enumerate([('FORM','PURPOSE','RULE')]+rows):
    for c,val in enumerate(row): set_cell(T.cell(r,c),val)
T=d.tables[9]
for r,(a,b) in enumerate(relrows): set_cell(T.cell(r,0),a); set_cell(T.cell(r,1),b)
T=d.tables[10]
set_cell(T.cell(0,0),'PROOF\nCLAIM / SCOPE / ENVIRONMENT / EVIDENCE / RESULT SHOULD accompany material conformance claims.')
set_cell(T.cell(0,1),'SNAPSHOT\nSPEC / lane A / lane B / MEMORY commit IDs MUST identify the reviewed release view.')
set_cell(T.cell(0,2),'VERSION\nTemplate version = FSL language. Project version = governed system. They MUST NOT be conflated.')
set_para(d.paragraphs[21],'PAGE 3 - TRACEABILITY / SNAPSHOT / CERTIFICATION')
set_para(d.paragraphs[22],'Use Page 3 only for complex projects. Three pages remains the maximum. This page records exact proof and release identity while implementation details remain in referenced supporting artifacts.')
set_para(d.paragraphs[23],'ARTICLE III - TRACEABILITY / AMENDMENT LAW')
set_para(d.paragraphs[24],'SCOPED CONFORMANCE EVIDENCE LEDGER')
set_para(d.paragraphs[25],'FOUR-SQUARE RELEASE SNAPSHOT')
set_para(d.paragraphs[26],'TEMPLATE / PROJECT VERSION GOVERNANCE')
set_para(d.paragraphs[27],'Every generated page MUST show Four-Square Template v1.3 separately from the governed project version. Template changes version the FSL language; ordinary project releases do not. Clause IDs MUST remain stable after publication, and detailed evidence SHOULD live outside the 1-3 page letterhead.')
T=d.tables[11]
new11=[
    ('CLAUSE','FORCE','TEXT'),
    ('LAW-TRACE-001','MUST','Every material binding clause MUST have a stable clause ID; a published ID MUST NOT be reused for a different meaning.'),
    ('LAW-VERSION-001','MUST NOT','FSL template version and governed project version MUST NOT be conflated.'),
    ('LAW-SNAPSHOT-001','MUST','An accepted Four-Square release MUST record the exact four lane commit IDs.'),
    ('LAW-AI-001','MUST NOT','An agent MUST NOT silently select a stronger or weaker interpretation when normative intent is ambiguous.')
]
for r,row in enumerate(new11):
    for c,val in enumerate(row): set_cell(T.cell(r,c),val)
T=d.tables[12]
headers=['CLAUSE ID','CLAIM','SCOPE','ENVIRONMENT','EVIDENCE','RESULT']
for c,h in enumerate(headers): set_cell(T.cell(0,c),h)
for r in range(1,len(T.rows)):
    vals=['[LAW-...]','[CLAIM]','[ARTIFACT / WORKFLOW]','[HOST / VERSION]','[TEST / LOG / HASH]','[CONFORMANT / EXCEPTION / NON-CONFORMANT / NOT YET PROVEN]']
    for c,v in enumerate(vals): set_cell(T.cell(r,c),v)
T=d.tables[13]
set_cell(T.cell(0,0),'FOUR SQUARES\nSPEC / lane A / lane B / MEMORY commit IDs MUST identify the reviewed release.')
set_cell(T.cell(0,1),'VERSION IDENTITY\nTemplate and project versions MUST remain separate.')
set_cell(T.cell(0,2),'SNAPSHOT\nThe snapshot MUST be declarative and MUST NOT replace Git as authority.')
set_cell(T.cell(0,3),'EVIDENCE\nClaims MUST remain scoped to what was actually proven.')
for s in d.sections:
    if s.footer.tables: set_cell(s.footer.tables[0].cell(0,1),'Traceability\nCLAUSE ID -> PROOF -> SNAPSHOT')
d.save(p3)

# Django sample
ps=OUT/'FSL-v1.3-Sample-Django-2P.docx'
shutil.copy2(SRC/'FSL-v1.2-Sample-Django-2P.docx',ps)
d=Document(ps); replace_version_everywhere(d, sample=True)
set_para(d.paragraphs[3],'Fixed lanes: SPEC + MEMORY. Middle lanes are CODE + DOCS for this illustrative sample. Framework/template version and sample/project version MUST remain separate identities.')
set_para(d.paragraphs[5],'CANONICAL WORKFLOWS / CONTRACTS / RELATIONS')
set_para(d.paragraphs[9],'PAGE 2 - ILLUSTRATIVE DJANGO TRACEABILITY')
set_para(d.paragraphs[10],'The clauses below demonstrate FSL v1.3 clause IDs, scoped proof, version separation and relation semantics. They are not upstream Django policy.')
set_para(d.paragraphs[11],'ARTICLE D1 - PUBLIC BEHAVIOR / TRACEABILITY')
set_para(d.paragraphs[12],'ARTICLE D2 - VERSION / AUTHORITY')
set_para(d.paragraphs[13],'REUSE / INHERIT / REFERENCE')
set_para(d.paragraphs[14],'Example review: [ ] clause IDs stable  [ ] code evidence  [ ] docs alignment  [ ] proof scope recorded  [ ] four commits snapshotted  [ ] active EXCEPTION listed.')
set_para(d.paragraphs[15],'ARTICLE D3 - SCOPED PROOF / CONFORMANCE')
T=d.tables[3]; set_cell(T.cell(0,1),'RELATION / EXECUTION SOURCE')
relations=['REUSE: URL routing + views + middleware','REUSE: ORM + migrations','REUSE: template engine','REUSE: forms + validators','REUSE: admin application','REFERENCE: Django test utilities + Python tooling']
for i,txt in enumerate(relations,1): set_cell(T.cell(i,1),txt)
set_cell(d.tables[5].cell(0,3),'MEMORY\nmemory-vN\nHuman+AI recall and examples. MEMORY may remember law; MEMORY cannot make law.')
set_cell(d.tables[6].cell(1,0),'RELEASE LAW - ONE VERSION, FOUR SQUARES, ALL GREEN\nA release MUST record the reviewed SPEC / CODE / DOCS / MEMORY commits and MUST NOT claim conformance beyond its scoped evidence.')
T=d.tables[7]
newrows=[
    ('CLAUSE ID','FORM','FORCE','CLAUSE'),
    ('LAW-DJANGO-001','LAW','MUST','A behavior declared as public MUST be supported by implementation evidence and corresponding developer guidance.'),
    ('GUARD-DJANGO-002','GUARDRAIL','MUST NOT','DOCS MUST NOT knowingly describe public behavior that the accepted CODE lane does not support.'),
    ('POLICY-DJANGO-003','POLICY','SHOULD','Compatibility and deprecation changes SHOULD record affected versions, migration guidance and evidence.'),
    ('CONTRACT-DJANGO-004','CONTRACT','MUST','A release contract MUST distinguish tested behavior from illustrative or provisional material.'),
    ('LAW-MEMORY-005','LAW','MUST NOT','MEMORY MUST NOT create new framework requirements unless promoted into SPEC.')
]
for r,row in enumerate(newrows):
    for c,val in enumerate(row): set_cell(T.cell(r,c),val)
T=d.tables[8]
set_cell(T.cell(0,0),'TRACEABILITY\nBinding clauses MUST use stable IDs and SHOULD map to evidence.')
set_cell(T.cell(0,1),'SNAPSHOT\nAccepted release view MUST record exact SPEC / CODE / DOCS / MEMORY commits.')
set_cell(T.cell(0,2),'AUTHORITY\nAgents MAY assist but MUST NOT invent law; MEMORY may remember law but cannot make law.')
T=d.tables[9]
while len(T.rows)<5: T.add_row()
proof=[
    ('CLAIM','Public behavior, implementation and developer guidance are aligned for the declared scope.'),
    ('SCOPE','Illustrative Django Four-Square sample; not an upstream governance claim.'),
    ('ENVIRONMENT','[DECLARED DJANGO / PYTHON / PLATFORM VERSIONS]'),
    ('EVIDENCE','[TEST / DOC / CHANGESET REFERENCES]'),
    ('RESULT','NOT YET PROVEN until real scoped evidence is attached; the sample MUST NOT imply upstream conformance.')
]
for r,(a,b) in enumerate(proof): set_cell(T.cell(r,0),a); set_cell(T.cell(r,1),b)
for s in d.sections:
    if s.footer.tables: set_cell(s.footer.tables[0].cell(0,1),'Traceability\nCLAUSE ID -> PROOF -> SNAPSHOT')
d.save(ps)

from docx.shared import RGBColor
WHITE=RGBColor(255,255,255)
def white_cell(cell,bold=True):
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.color.rgb=WHITE
            if bold: r.bold=True
for path in OUT.glob('*.docx'):
    dd=Document(path)
    if len(dd.tables)>3: white_cell(dd.tables[3].cell(0,1))
    if 'Sample-Django' in path.name:
        if len(dd.tables)>6: white_cell(dd.tables[6].cell(1,0),bold=False)
        if len(dd.tables)>7:
            for c in range(len(dd.tables[7].columns)): white_cell(dd.tables[7].cell(0,c))
    else:
        if len(dd.tables)>7: white_cell(dd.tables[7].cell(1,0),bold=False)
        if len(dd.tables)>8:
            for c in range(len(dd.tables[8].columns)): white_cell(dd.tables[8].cell(0,c))
        if len(dd.tables)>11:
            for c in range(len(dd.tables[11].columns)): white_cell(dd.tables[11].cell(0,c))
        if len(dd.tables)>12:
            for c in range(len(dd.tables[12].columns)): white_cell(dd.tables[12].cell(0,c))
    dd.save(path)
print('FSL v1.3 DOCX build complete')
