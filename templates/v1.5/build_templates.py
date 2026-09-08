#!/usr/bin/env python3
"""Build the editable FSL v1.5 letterheads and illustrative Django sample.

No older FSL generator, repository checkout, YAML or network access is needed.
Output PDFs are rendered from these DOCX files, not authored separately.
"""
from __future__ import annotations

import argparse
import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

VERSION = "1.5"
FOOTER_AXIOM = "FSL owns the method. The project owns the work."
FONT = "Arial"
INK, NAVY, MUTED, LINE = "172B3A", "183B54", "526474", "C9D3DB"
PALE, WHITE = "F2F5F7", "FFFFFF"
TERMS = (
    "CONFORMANT WITH EXCEPTION", "NOT YET PROVEN", "NON-CONFORMANT",
    "MUST NOT", "SHOULD NOT", "SHALL NOT", "MUST", "SHOULD", "SHALL", "MAY",
    "CONFORMANT", "EXCEPTION", "WAIVER",
)
NORM = re.compile(r"(?<![A-Za-z0-9_-])(" + "|".join(map(re.escape, TERMS)) + r")(?![A-Za-z0-9_-])")
WIDTH = 510  # points; matches the page's writable width
DOCS = {
    "D1": ("Django 5.2 overview", "https://docs.djangoproject.com/en/5.2/intro/overview/"),
    "D2": ("Django 5.2 testing", "https://docs.djangoproject.com/en/5.2/topics/testing/overview/"),
    "D3": ("Django 5.2 shortcuts", "https://docs.djangoproject.com/en/5.2/topics/http/shortcuts/"),
}


def font(run, size=9, bold=False, color=INK, underline=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.underline = underline
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    return run


def rich(p, text: str, size=9, bold=False, color=INK):
    """Emphasize only complete normative terms, preserving clause identifiers."""
    for chunk in NORM.split(text):
        if not chunk:
            continue
        normative = chunk in TERMS
        font(p.add_run(chunk), size, bold or normative, color, normative)
    return p


def paragraph(parent, text="", size=9, before=0, after=3, bold=False, color=INK,
              keep=False, align=None):
    p = parent.add_paragraph()
    f = p.paragraph_format
    f.space_before, f.space_after = Pt(before), Pt(after)
    f.line_spacing = 1.07
    f.keep_together = True
    f.keep_with_next = keep
    if align is not None:
        p.alignment = align
    rich(p, text, size, bold, color)
    return p


def label(d, text, before=8):
    p = paragraph(d, text, size=9.5, before=before, after=4, bold=True, color=NAVY, keep=True)
    p.style = d.styles["Heading 2"]
    return p


def shade(c, fill):
    pr = c._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:fill"), fill)
    pr.append(el)


def borders(t, color=LINE, horizontal=True, vertical=False, outer=False):
    pr = t._tbl.tblPr
    el = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement("w:" + edge)
        enabled = ((horizontal and edge in ("top", "bottom", "insideH")) or
                   (vertical and edge == "insideV") or outer)
        b.set(qn("w:val"), "single" if enabled else "nil")
        b.set(qn("w:sz"), "5")
        b.set(qn("w:color"), color)
        el.append(b)
    pr.append(el)


def table(d, widths, rows, header=None, size=8.7):
    t = d.add_table(rows=0, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    pr = t._tbl.tblPr
    w = pr.find(qn("w:tblW")); w.set(qn("w:w"), str(int(sum(widths)*20))); w.set(qn("w:type"), "dxa")
    for col, width in zip(t.columns, widths):
        col.width = Pt(width)
    data = ([header] if header else []) + list(rows)
    for ri, values in enumerate(data):
        row = t.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if header and ri == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for ci, value in enumerate(values):
            c = row.cells[ci]; c.width = Pt(widths[ci])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            tcpr = c._tc.get_or_add_tcPr()
            mar = OxmlElement("w:tcMar")
            for edge, val in (("top",70),("bottom",70),("start",100),("end",100)):
                e = OxmlElement("w:"+edge); e.set(qn("w:w"), str(val)); e.set(qn("w:type"),"dxa"); mar.append(e)
            tcpr.append(mar)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            p.paragraph_format.keep_together = True
            rich(p, value, 7.9 if header and ri==0 else size,
                 bold=bool(header and ri==0), color=NAVY if header and ri==0 else INK)
            if header and ri==0:
                shade(c, PALE)
    borders(t)
    return t


def field(p, code):
    el=OxmlElement("w:fldSimple"); el.set(qn("w:instr"), code)
    r=OxmlElement("w:r"); rp=OxmlElement("w:rPr")
    f=OxmlElement("w:rFonts"); f.set(qn("w:ascii"),FONT); f.set(qn("w:hAnsi"),FONT); rp.append(f)
    sz=OxmlElement("w:sz"); sz.set(qn("w:val"),"15"); rp.append(sz); r.append(rp)
    tx=OxmlElement("w:t"); tx.text="1"; r.append(tx); el.append(r); p._p.append(el)


def link(p, text, url, size=8.5):
    hl=OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    r=OxmlElement("w:r"); pr=OxmlElement("w:rPr")
    fs=OxmlElement("w:rFonts"); fs.set(qn("w:ascii"),FONT); fs.set(qn("w:hAnsi"),FONT); pr.append(fs)
    sz=OxmlElement("w:sz"); sz.set(qn("w:val"),str(int(size*2))); pr.append(sz)
    col=OxmlElement("w:color"); col.set(qn("w:val"),NAVY); pr.append(col)
    u=OxmlElement("w:u"); u.set(qn("w:val"),"single"); pr.append(u); r.append(pr)
    tx=OxmlElement("w:t"); tx.text=text; r.append(tx); hl.append(r); p._p.append(hl)


def new_document(sample=False):
    d=Document(); s=d.sections[0]
    s.page_width, s.page_height = Pt(595.28), Pt(841.89)
    s.left_margin = s.right_margin = Pt((595.28-WIDTH)/2)
    s.top_margin, s.bottom_margin = Pt(92), Pt(59)
    s.header_distance, s.footer_distance = Pt(26), Pt(22)
    normal=d.styles["Normal"]; normal.font.name=FONT; normal.font.size=Pt(9)
    normal.paragraph_format.space_after=Pt(3)
    normal.paragraph_format.line_spacing=1.07
    for style in ("Heading 1", "Heading 2"):
        d.styles[style].font.name=FONT
        d.styles[style].font.color.rgb=RGBColor.from_string(NAVY)
        d.styles[style].paragraph_format.keep_with_next=True
    d.core_properties.title = f"FSL v{VERSION} - {'Django illustrative sample' if sample else 'Project letterhead template'}"
    d.core_properties.subject = "Compact project constitution and development method"
    d.core_properties.author="FSL"
    d.core_properties.last_modified_by="FSL"
    d.core_properties.created = d.core_properties.modified = datetime(2026,9,8,0,0,0,tzinfo=timezone.utc)
    h=s.header
    p=h.paragraphs[0]; p.style=d.styles["Normal"]; p.paragraph_format.space_after=Pt(4)
    p.paragraph_format.tab_stops.add_tab_stop(Pt(WIDTH), WD_TAB_ALIGNMENT.RIGHT)
    font(p.add_run("FSL  /  FOUR-SQUARE LETTERHEAD"), 8, True, NAVY)
    font(p.add_run("\t" + ("Sample edition v1.5" if sample else "Project version [vN]")), 8, True, MUTED)
    p=paragraph(h,"DJANGO" if sample else "[PROJECT / PRODUCT]",size=21,after=2,bold=True,color=NAVY)
    repo = "ILLUSTRATIVE SAMPLE  •  Upstream reference: django/django" if sample else "Project repository: [owner/repository]  •  Owner: [responsible person/team]"
    p=paragraph(h,repo,size=8,after=5,color=MUTED)
    pr=p._p.get_or_add_pPr(); bd=OxmlElement("w:pBdr"); b=OxmlElement("w:bottom")
    b.set(qn("w:val"),"single"); b.set(qn("w:sz"),"12"); b.set(qn("w:color"),NAVY); b.set(qn("w:space"),"6"); bd.append(b); pr.append(bd)
    footer=s.footer
    p=footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(4)
    font(p.add_run(FOOTER_AXIOM),8.3,True,NAVY)
    p=paragraph(footer,"",size=7.5,after=0,color=MUTED)
    p.paragraph_format.tab_stops.add_tab_stop(Pt(WIDTH), WD_TAB_ALIGNMENT.RIGHT)
    font(p.add_run(f"Four-Square Template v{VERSION}"),7.5,False,MUTED)
    font(p.add_run("\tPage "),7.5,False,MUTED); field(p,"PAGE")
    font(p.add_run(" / "),7.5,False,MUTED); field(p,"NUMPAGES")
    return d


def squares(d, sample=False):
    a,b=("CODE","DOCS") if sample else ("[LANE A]","[LANE B]")
    va,vb=("code","docs") if sample else ("lane-a","lane-b")
    v="N"
    data=[
        ("01  SPEC",f"spec-v{v}","Owns binding intent, scope and contracts.",
         "Gate: stable clauses + scoped acceptance criteria."),
        (f"02  {a}",f"{va}-v{v}","Owns the primary implementation or deliverable." if not sample else "Owns framework code and its executable checks.",
         "Gate: the declared deliverable passes its checks."),
        (f"03  {b}",f"{vb}-v{v}","Owns the second durable project concern." if not sample else "Owns developer guidance aligned with code.",
         "Gate: usable, aligned and independently checked."),
        ("04  MEMORY",f"memory-v{v}","Owns recall, decisions, examples and evidence links.",
         "Gate: understandable; no new law outside SPEC."),
    ]
    t=table(d,[255,255],[["",""],["",""]],size=8.5)
    borders(t,outer=True)
    for i,(title,branch,desc,gate) in enumerate(data):
        c=t.cell(i//2,i%2); shade(c, PALE if i in (0,3) else WHITE)
        p=c.paragraphs[0]
        p.paragraph_format.space_after=Pt(4)
        font(p.add_run(title),10,True,NAVY)
        font(p.add_run("  /  "+branch),8,False,MUTED)
        paragraph(c,desc,size=8.7,after=3)
        paragraph(c,gate,size=8.2,after=1,color=MUTED)
    return t

def page_one(d, sample=False):
    if sample:
        paragraph(d,"A teaching overlay, not upstream Django policy or a claim of tested behavior.",size=8.3,color=MUTED,after=4)
        paragraph(d,"PURPOSE  Database-backed web applications using models, URLs, views and templates. [D1]",size=9,after=3)
        paragraph(d,"BOUNDARY  The sample governs no upstream project. Deployment, business rules and production proof remain outside this illustration.",size=9,after=3)
    else:
        paragraph(d,"PURPOSE  [Who this serves, the outcome it promises, and its distinguishing principle.]",size=9,after=3)
        paragraph(d,"BOUNDARY  [What this contract includes; what it deliberately does not own.]",size=9,after=3)
    label(d,"FOUR SQUARES",before=7)
    squares(d,sample)
    paragraph(d,"SPEC + MEMORY are fixed; the two middle lanes are project-defined. Branch suffix vN follows the project, not the FSL template version.",size=7.8,color=MUTED,before=3,after=1)
    label(d,"BINDING CONTRACT",before=7)
    paragraph(d,"MUST / MUST NOT = binding   •   SHOULD / SHOULD NOT = guidance   •   MAY = permission",size=7.8,after=4)
    if sample:
        rows=[
            ("LAW-DJANGO-001","Public behavior MUST have scoped code evidence and aligned guidance.","Code test + relevant guidance"),
            ("GUARD-DJANGO-002","DOCS MUST NOT describe behavior unsupported by accepted CODE.","Documentation review"),
            ("CONTRACT-DJANGO-004","Release evidence MUST separate tested behavior from illustration and state scope.","Scoped release review"),
        ]
    else:
        rows=[
            ("LAW-[SCOPE]-001","[Subject] MUST [observable requirement].","[Check / evidence path]"),
            ("CONTRACT-[SCOPE]-002","[Interface] MUST [observable promise].","[Acceptance evidence]"),
            ("GUARD-[SCOPE]-003","[Operation] MUST NOT [forbidden effect].","[Negative / boundary test]"),
        ]
    table(d,[119,277,114],rows,["STABLE CLAUSE ID","REQUIREMENT / PROMISE","VERIFICATION"],8.5)
    paragraph(d,"SPEC governs. Binding requirements MUST NOT be silently weakened; missing evidence is NOT YET PROVEN.",size=8.2,before=3,after=1)
    label(d,"DEV STRATEGY  /  INITIAL ITERATION",before=7)
    initial="I-001: one request → URL → view → one stored record → HTML response → prove → grow." if sample else "I-001: tiny [input] → [required touchpoints] → [observable output] → prove → grow."
    paragraph(d,initial,size=8.6,bold=True,after=3)
    paragraph(d,"Parallel E2E / E2M / M2E MAY be used. M = named midpoint contract. Partial or component proof is not E2E. No full iteration plan is required.",size=8.0,color=MUTED,after=1)
    label(d,"TECHNOLOGIES, LIBRARIES & REFERENCES",before=7)
    if sample:
        rows=[("Django facilities","SHOULD REUSE for the initial slice","[D1] Overview"),
              ("Django test tools","SHOULD REUSE for scoped checks","[D2] Testing"),
              ("Django shortcuts","MAY REUSE for view responses","[D3] Shortcuts")]
    else:
        rows=[("[Technology / library]","[MUST / SHOULD / MAY] [purpose]","[Canonical link / source ID]"),
              ("[Existing capability]","[REUSE / INHERIT / REFERENCE]","[Source; pin only if needed]")]
    table(d,[126,242,142],rows,["ITEM","INTENT / USE","SOURCE"],8.2)
    paragraph(d,"Only essential choices belong here. A category or link alone creates no obligation; a binding choice needs a SPEC clause ID.",size=7.8,color=MUTED,before=3,after=1)
    label(d,"RELEASE CHECK  /  ONE VERSION, FOUR SQUARES, ALL GREEN",before=7)
    a,b=("CODE","DOCS") if sample else ("LANE A","LANE B")
    paragraph(d,f"[ ] SPEC stable     [ ] {a} verified     [ ] {b} usable and aligned     [ ] MEMORY understandable",size=8.0,after=3)
    paragraph(d,"Status: NOT YET PROVEN  |  Evidence / exact four-commit snapshot: " + ("not supplied; example only." if sample else "[project-local path]."),size=8.1,after=3)
    paragraph(d,"main1 = this repository's first commit, fixed. Four lane branches vN → matching lanes vN+1. main = accepted integration after gates.",size=7.6,color=MUTED,before=2,after=0)


def next_page(d, title, sub):
    d.add_page_break()
    p=paragraph(d,title,size=16,bold=True,color=NAVY,after=6); p.style=d.styles["Heading 1"]
    paragraph(d,sub,size=9,color=MUTED,after=9)


def page_two(d):
    next_page(d,"CONTRACTS & IMPLEMENTATION NOTES", "Use this page only when essential detail will not fit on Page 1. Delete unused sections.")
    label(d,"INTERFACES & OBSERVABLE OUTCOMES",before=3)
    table(d,[113,224,173],[
        ("[CONTRACT-ID]","[Input / caller → interface → consumer]","[Result and how it is checked]"),
        ("[CONTRACT-ID]","[State, event or artifact exchanged]","[Failure response / recovery promise]"),
        ("[GUARD-ID]","[Permission or destructive boundary]","[Check that the boundary holds]"),
    ],["CLAUSE","BOUNDARY / INTERFACE","OBSERVABLE PROMISE"],9)
    paragraph(d,"Use a named interface or observable contract for a midpoint. E2M and M2E are partial paths until their integrated E2E outcome is demonstrated.",size=9,before=5,after=5)
    label(d,"REUSE WHAT IS ALREADY PROVEN")
    table(d,[91,419],[
        ("REUSE","Use a capability. Do not automatically adopt the source project's rules or assumptions."),
        ("INHERIT","Explicitly adopt a requirement or contract into project SPEC; identify its source and scope."),
        ("REFERENCE","Consult knowledge or evidence without adopting its normative authority."),
    ],size=9)
    label(d,"ESSENTIAL IMPLEMENTATION REFERENCES")
    table(d,[118,229,163],[
        ("[Runtime / library]","[Use + reason + MUST/SHOULD/MAY]","[Canonical source / clause ID]"),
        ("[Existing implementation]","[REUSE + scope of the capability]","[Repository / commit, if material]"),
        ("[External contract]","[INHERIT only where explicitly adopted]","[Standard / version, if material]"),
    ],["ITEM","INTENT / PURPOSE","SOURCE / OPTIONAL PIN"],9)
    paragraph(d,"Add a version, date or commit only where it changes the meaning. Keep transitive dependencies, large inventories and moving test matrices in referenced project files.",size=8.8,before=5,after=3,color=MUTED)
    label(d,"SCOPED PROOF")
    paragraph(d,"Claim: [what is being asserted]. Scope: [artifact + workflow]. Environment: [platform + relevant versions]. Evidence: [test, review or log]. Result: NOT YET PROVEN until checked.",size=9,after=5)
    paragraph(d,"Fixtures or stubs MAY support a declared partial check. They MUST NOT be presented as proof of a real integration they replace.",size=9,after=3)
    label(d,"OPEN QUESTIONS & EXCEPTIONS")
    table(d,[112,398],[
        ("Unresolved intent","[Question + person responsible for deciding it]. Agents do not invent the missing rule."),
        ("Active exception","[Clause ID + scope + authority + reason + evidence + review/expiry]; otherwise state NONE."),
    ],size=9)
    paragraph(d,"A recorded exception does not silently change the original requirement. Conformance remains limited to the declared scope and authorized exception.",size=8.8,before=5,color=MUTED)


def page_three(d):
    next_page(d,"RELEASE & EVIDENCE", "A reviewed release view, not an iteration tracker. Keep full logs and implementation detail in project-local supporting files.")
    label(d,"EXACT FOUR-SQUARE SNAPSHOT",before=3)
    paragraph(d,"Release: [project release ID]   •   Project repository: [owner/repository]   •   FSL template: v1.5",size=9,after=6)
    table(d,[75,125,178,132],[
        ("SPEC","spec-vN","[Full reviewed commit ID]","[Gate evidence]"),
        ("[LANE A]","[lane-a]-vN","[Full reviewed commit ID]","[Gate evidence]"),
        ("[LANE B]","[lane-b]-vN","[Full reviewed commit ID]","[Gate evidence]"),
        ("MEMORY","memory-vN","[Full reviewed commit ID]","[Gate evidence]"),
    ],["SQUARE","PROJECT-LOCAL BRANCH","REVIEWED COMMIT","GATE"],8.8)
    paragraph(d,"Integration: [full accepted commit ID]. Publication: [branch / tag / release URL and actual state]. A branch name alone is not an immutable snapshot.",size=9,before=5,after=3)
    paragraph(d,"LAW-SNAPSHOT-001 / 002: record the four reviewed commits; the snapshot does not replace Git. Do not invent a commit or claim publication from a local package.",size=8.5,color=MUTED,after=3)
    label(d,"SCOPED ACCEPTANCE EVIDENCE")
    table(d,[112,142,132,124],[
        ("[CLAUSE-ID]","[Claim + artifact / path]","[Environment + evidence link]","NOT YET PROVEN"),
        ("[CLAUSE-ID]","[Second material promise]","[Environment + evidence link]","NOT YET PROVEN"),
    ],["CLAUSE","CLAIM / SCOPE","PROOF CONTEXT","RESULT"],8.8)
    paragraph(d,"A changed component can pass its own gate without establishing the complete user path. Scope the statement to the proof actually obtained.",size=9,before=5,after=3)
    label(d,"CONFORMANCE DECLARATION")
    table(d,[166,344],[
        ("CONFORMANT","Applicable binding requirements have evidence; no unresolved binding violation."),
        ("CONFORMANT WITH EXCEPTION","An authorized, explicit and scoped exception accompanies otherwise supported claims."),
        ("NON-CONFORMANT","A binding requirement is known to be violated."),
        ("NOT YET PROVEN","Evidence or a necessary interpretation is unresolved."),
    ],size=9)
    paragraph(d,"Declared state: [NOT YET PROVEN]   •   Reviewer / date: [owner / date]   •   Active exceptions: [NONE / IDs]",size=8.7,before=6,after=3)
    label(d,"AMENDMENT & HANDOVER")
    paragraph(d,"Change: [clause IDs / project version]. Compatibility or migration impact: [brief impact or NONE]. Supporting decisions and evidence: [MEMORY paths].",size=9,after=5)
    paragraph(d,"Change binding intent in SPEC; align both project lanes and MEMORY. Preserve published clause IDs and the frozen release view. MEMORY may remember law; MEMORY cannot make law.",size=9,after=3)


def sample_page_two(d):
    next_page(d,"DJANGO / SCOPED ILLUSTRATION", "Sample edition v1.5. Reference baseline: Django 5.2 documentation, not an asserted upstream release or executed test environment.")
    label(d,"ONE SMALL PATH; A CLEAR CLAIM",before=3)
    paragraph(d,"Illustrative intent: expose one stored article through one URL and return its title in an HTML response. Django's overview demonstrates URL routing, views, model access and template rendering. [D1]",size=9.2,after=6)
    table(d,[120,226,164],[
        ("URL + view","Route the request to the article view.","Correct view and arguments [D1]"),
        ("Model + database","Read one known article record.","Expected stored title [D1]"),
        ("Template + response","Render the title into the returned HTML.","Expected response content [D1, D3]"),
    ],["TOUCHPOINTS","TINY BEHAVIOR","PLANNED CHECK"],9)
    paragraph(d,"This is an example design, not execution evidence. The chosen claim ends at the HTML response. Browser rendering, JavaScript, hosting and production integration remain unproven.",size=9,before=6,after=3)
    label(d,"MIDPOINT & PARALLEL WORK")
    paragraph(d,"M = the view's response contract: known request, expected status and article title. The request-to-view side and the view-to-HTML side MAY develop separately. Neither side alone proves their integrated path.",size=9.2,after=5)
    paragraph(d,"An unrelated second E2E path MAY proceed in parallel. Keep its proof separate; passing one path does not establish another. No later-iteration schedule is prescribed here.",size=9.2,after=3)
    label(d,"REFERENCE INTENT")
    table(d,[128,382],[
        ("Django facilities [D1]","REUSE for models, URL handling, views and templates in this example. This does not import an upstream governance policy."),
        ("Testing [D2]","REFERENCE the documented test runner and database-aware tests when writing checks. A planned test is not a passed test."),
        ("Shortcuts [D3]","REFERENCE the documented response helpers; using a helper does not establish browser or deployment proof."),
    ],size=9)
    label(d,"PROOF RECORD / CURRENT STATE")
    paragraph(d,"Claim: one stored article is returned in HTML. Scope: the illustrative request-to-response path. Environment: not executed. Evidence: no runtime test supplied. Result: NOT YET PROVEN.",size=9.2,after=5)
    paragraph(d,"Release snapshot: none. No SPEC / CODE / DOCS / MEMORY commits or upstream release status are asserted. The sample's own page-layout checks are separate from product conformance.",size=9,color=MUTED,after=3)
    label(d,"CANONICAL SOURCES")
    for key,(title,url) in DOCS.items():
        p=paragraph(d,"",size=8.5,after=4)
        font(p.add_run(f"[{key}] "),8.5,True,NAVY)
        link(p,title,url,8.5)
        font(p.add_run("  ·  "+url.replace("https://", "")),7.5,False,MUTED)
    paragraph(d,"Source baseline consulted 8 September 2026. The proposed workflow and all binding sample clauses are FSL-authored illustration, not Django's own policy.",size=8.3,color=MUTED,before=3,after=0)


def build(out: Path):
    out.mkdir(parents=True, exist_ok=True)
    for pages in (1,2,3):
        d=new_document()
        page_one(d)
        if pages>=2: page_two(d)
        if pages>=3: page_three(d)
        path=out/f"FSL-v{VERSION}-Template-{pages}P.docx"
        d.save(path)
        print(path.name)
    d=new_document(sample=True); page_one(d,sample=True); sample_page_two(d)
    path=out/f"FSL-v{VERSION}-Sample-Django-2P.docx"; d.save(path); print(path.name)


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output",type=Path,default=Path(__file__).resolve().parents[1])
    args=parser.parse_args()
    build(args.output)

if __name__=="__main__":
    main()
