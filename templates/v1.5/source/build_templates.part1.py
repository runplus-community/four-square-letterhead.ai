#!/usr/bin/env python3
"""Build the editable FSL v1.5 letterheads and illustrative Django sample.

No older FSL generator, repository checkout, YAML or network access is needed.
Output PDFs are rendered from these DOCX files, not authored separately.
"""
from __future__ import annotations

import argparse
import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

VERSION = "1.5"
FOOTER_AXIOM = "FSL owns the method. The project owns the work."
FONT = "Arial"
INK, NAVY, MUTED, LINE = "172B3A", "183B54", "526474", "C9D3DB"
PALE, WHITE = "F2F5F7", "FFFFFF"
TERMS = (
    "CONFORMANT WITH EXCEPTION", "NOT YET PROVEN", "NON-CONFORMANT",
    "MUST NOT", "SHOULD NOT", "SHALL NOT", "MUST", "SHOULD", "SHALL", "MAY",
    "CONFORMANT", "EXCEPTION", "WAIVER",
)
NORM = re.compile(r"(?<![A-Za-z0-9_-])(" + "|".join(map(re.escape, TERMS)) + r")(?![A-Za-z0-9_-])")
WIDTH = 510  # points; matches the page's writable width
DOCS = {
    "D1": ("Django 5.2 overview", "https://docs.djangoproject.com/en/5.2/intro/overview/"),
    "D2": ("Django 5.2 testing", "https://docs.djangoproject.com/en/5.2/topics/testing/overview/"),
    "D3": ("Django 5.2 shortcuts", "https://docs.djangoproject.com/en/5.2/topics/http/shortcuts/"),
}


def font(run, size=9, bold=False, color=INK, underline=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.underline = underline
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    return run


def rich(p, text: str, size=9, bold=False, color=INK):
    """Emphasize only complete normative terms, preserving clause identifiers."""
    for chunk in NORM.split(text):
        if not chunk:
            continue
        normative = chunk in TERMS
        font(p.add_run(chunk), size, bold or normative, color, normative)
    return p


def paragraph(parent, text="", size=9, before=0, after=3, bold=False, color=INK,
              keep=False, align=None):
    p = parent.add_paragraph()
    f = p.paragraph_format
    f.space_before, f.space_after = Pt(before), Pt(after)
    f.line_spacing = 1.07
    f.keep_together = True
    f.keep_with_next = keep
    if align is not None:
        p.alignment = align
    rich(p, text, size, bold, color)
    return p


def label(d, text, before=8):
    p = paragraph(d, text, size=9.5, before=before, after=4, bold=True, color=NAVY, keep=True)
    p.style = d.styles["Heading 2"]
    return p


def shade(c, fill):
    pr = c._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:fill"), fill)
    pr.append(el)


def borders(t, color=LINE, horizontal=True, vertical=False, outer=False):
    pr = t._tbl.tblPr
    el = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement("w:" + edge)
        enabled = ((horizontal and edge in ("top", "bottom", "insideH")) or
                   (vertical and edge == "insideV") or outer)
        b.set(qn("w:val"), "single" if enabled else "nil")
        b.set(qn("w:sz"), "5")
        b.set(qn("w:color"), color)
        el.append(b)
    pr.append(el)


def table(d, widths, rows, header=None, size=8.7):
    t = d.add_table(rows=0, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    pr = t._tbl.tblPr
    w = pr.find(qn("w:tblW")); w.set(qn("w:w"), str(int(sum(widths)*20))); w.set(qn("w:type"), "dxa")
    for col, width in zip(t.columns, widths):
        col.width = Pt(width)
    data = ([header] if header else []) + list(rows)
    for ri, values in enumerate(data):
        row = t.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if header and ri == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for ci, value in enumerate(values):
            c = row.cells[ci]; c.width = Pt(widths[ci])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            tcpr = c._tc.get_or_add_tcPr()
            mar = OxmlElement("w:tcMar")
            for edge, val in (("top",70),("bottom",70),("start",100),("end",100)):
                e = OxmlElement("w:"+edge); e.set(qn("w:w"), str(val)); e.set(qn("w:type"),"dxa"); mar.append(e)
            tcpr.append(mar)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            p.paragraph_format.keep_together = True
            rich(p, value, 7.9 if header and ri==0 else size,
                 bold=bool(header and ri==0), color=NAVY if header and ri==0 else INK)
            if header and ri==0:
                shade(c, PALE)
    borders(t)
    return t


def field(p, code):
    el=OxmlElement("w:fldSimple"); el.set(qn("w:instr"), code)
    r=OxmlElement("w:r"); rp=OxmlElement("w:rPr")
    f=OxmlElement("w:rFonts"); f.set(qn("w:ascii"),FONT); f.set(qn("w:hAnsi"),FONT); rp.append(f)
    sz=OxmlElement("w:sz"); sz.set(qn("w:val"),"15"); rp.append(sz); r.append(rp)
    tx=OxmlElement("w:t"); tx.text="1"; r.append(tx); el.append(r); p._p.append(el)


def link(p, text, url, size=8.5):
    hl=OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    r=OxmlElement("w:r"); pr=OxmlElement("w:rPr")
    fs=OxmlElement("w:rFonts"); fs.set(qn("w:ascii"),FONT); fs.set(qn("w:hAnsi"),FONT); pr.append(fs)
    sz=OxmlElement("w:sz"); sz.set(qn("w:val"),str(int(size*2))); pr.append(sz)
    col=OxmlElement("w:color"); col.set(qn("w:val"),NAVY); pr.append(col)
    u=OxmlElement("w:u"); u.set(qn("w:val"),"single"); pr.append(u); r.append(pr)
    tx=OxmlElement("w:t"); tx.text=text; r.append(tx); hl.append(r); p._p.append(hl)


def new_document(sample=False):
    d=Document(); s=d.sections[0]
    s.page_width, s.page_height = Pt(595.28), Pt(841.89)
    s.left_margin = s.right_margin = Pt((595.28-WIDTH)/2)
    s.top_margin, s.bottom_margin = Pt(92), Pt(59)
    s.header_distance, s.footer_distance = Pt(26), Pt(22)
    normal=d.styles["Normal"]; normal.font.name=FONT; normal.font.size=Pt(9)
    normal.paragraph_format.space_after=Pt(3)
    normal.paragraph_format.line_spacing=1.07
    for style in ("Heading 1", "Heading 2"):
        d.styles[style].font.name=FONT
        d.styles[style].font.color.rgb=RGBColor.from_string(NAVY)
        d.styles[style].paragraph_format.keep_with_next=True
    d.core_properties.title = f"FSL v{VERSION} - {'Django illustrative sample' if sample else 'Project letterhead template'}"
    d.core_properties.subject = "Compact project constitution and development method"
    d.core_properties.author="FSL"
    d.core_properties.last_modified_by="FSL"
    d.core_properties.created = d.core_properties.modified = datetime(2026,9,8,0,0,0,tzinfo=timezone.utc)
    h=s.header
    p=h.paragraphs[0]; p.style=d.styles["Normal"]; p.paragraph_format.space_after=Pt(4)
    p.paragraph_format.tab_stops.add_tab_stop(Pt(WIDTH), WD_TAB_ALIGNMENT.RIGHT)
    font(p.add_run("FSL  /  FOUR-SQUARE LETTERHEAD"), 8, True, NAVY)
    font(p.add_run("\t" + ("Sample edition v1.5" if sample else "Project version [vN]")), 8, True, MUTED)
    p=paragraph(h,"DJANGO" if sample else "[PROJECT / PRODUCT]",size=21,after=2,bold=True,color=NAVY)
    repo = "ILLUSTRATIVE SAMPLE  •  Upstream reference: django/django" if sample else "Project repository: [owner/repository]  •  Owner: [responsible person/team]"
    p=paragraph(h,repo,size=8,after=5,color=MUTED)
    pr=p._p.get_or_add_pPr(); bd=OxmlElement("w:pBdr"); b=OxmlElement("w:bottom")
    b.set(qn("w:val"),"single"); b.set(qn("w:sz"),"12"); b.set(qn("w:color"),NAVY); b.set(qn("w:space"),"6"); bd.append(b); pr.append(bd)
    footer=s.footer
    p=footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(4)
    font(p.add_run(FOOTER_AXIOM),8.3,True,NAVY)
    p=paragraph(footer,"",size=7.5,after=0,color=MUTED)
    p.paragraph_format.tab_stops.add_tab_stop(Pt(WIDTH), WD_TAB_ALIGNMENT.RIGHT)
    font(p.add_run(f"Four-Square Template v{VERSION}"),7.5,False,MUTED)
    font(p.add_run("\tPage "),7.5,False,MUTED); field(p,"PAGE")
    font(p.add_run(" / "),7.5,False,MUTED); field(p,"NUMPAGES")
    return d


def squares(d, sample=False):
    a,b=("CODE","DOCS") if sample else ("[LANE A]","[LANE B]")
    va,vb=("code","docs") if sample else ("lane-a","lane-b")
    v="N"
    data=[
        ("01  SPEC",f"spec-v{v}","Owns binding intent, scope and contracts.",
         "Gate: stable clauses + scoped acceptance criteria."),
        (f"02  {a}",f"{va}-v{v}","Owns the primary implementation or deliverable." if not sample else "Owns framework code and its executable checks.",
         "Gate: the declared deliverable passes its checks."),
        (f"03  {b}",f"{vb}-v{v}","Owns the second durable project concern." if not sample else "Owns developer guidance aligned with code.",
         "Gate: usable, aligned and independently checked."),
        ("04  MEMORY",f"memory-v{v}","Owns recall, decisions, examples and evidence links.",
         "Gate: understandable; no new law outside SPEC."),
    ]
    t=table(d,[255,255],[["",""],["",""]],size=8.5)
    borders(t,outer=True)
    for i,(title,branch,desc,gate) in enumerate(data):
        c=t.cell(i//2,i%2); shade(c, PALE if i in (0,3) else WHITE)
        p=c.paragraphs[0]
        p.paragraph_format.space_after=Pt(4)
        font(p.add_run(title),10,True,NAVY)
        font(p.add_run("  /  "+branch),8,False,MUTED)
        paragraph(c,desc,size=8.7,after=3)
        paragraph(c,gate,size=8.2,after=1,color=MUTED)
    return t
