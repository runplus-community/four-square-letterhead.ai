#!/usr/bin/env python3
"""Verify the v1.8 document contract, exact method text and optional package integrity."""
import argparse, hashlib, re, tempfile
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
import fitz
from build import AXIOM, NORM, blocks
DOCS={'Plain-English-Method-3P':(3,0),'Template-1P':(1,1),'Sample-Django-1P':(1,1)}
BODY_HASH={'Template-1P':'6283e74c17d0efdebb3a42113b3536ec0840b4cc8f2e729110ce7e0260ba149d','Sample-Django-1P':'827d62a23c4ed593ab97599f1bcef9db996fddb89cdabfe4c30eca2c4d5b1c89'}
def require(ok,message):
    if not ok: raise ValueError(message)
def paragraphs(doc):
    for el in [doc._element,*[s.header._element for s in doc.sections],*[s.footer._element for s in doc.sections]]:
        yield from el.xpath('.//w:p')
def check_docx(path,method,label):
    d=Document(path)
    require(len(d.tables)==DOCS[label][1],'Unexpected table: '+path.name)
    text='\n'.join(''.join(p.xpath('.//w:t/text()')) for p in paragraphs(d))
    for token in ('Four-Square Template v1.8',AXIOM,'R-01'): require(token in text,'Missing identity or route: '+token)
    require('LOCAL CANDIDATE' not in text and 'I-001' not in text,'Stale status or compulsory iteration')
    count=0
    for p in paragraphs(d):
        rr=p.xpath('.//w:r');tt=[''.join(r.xpath('.//w:t/text()')) for r in rr]
        for match in NORM.finditer(''.join(tt)):
            count+=1;offset=0
            for r,t in zip(rr,tt):
                end=offset+len(t)
                if match.start()<end and offset<match.end():
                    b=r.find('w:rPr/w:b',r.nsmap);u=r.find('w:rPr/w:u',r.nsmap)
                    require(b is not None and b.get(qn('w:val')) not in ('0','false'),'Missing bold: '+match.group())
                    require(u is not None and u.get(qn('w:val')) not in ('0','none','false'),'Missing underline: '+match.group())
                offset=end
    require(count>=5,'Too few normative markers')
    if label.startswith('Plain-English'):
        expected=[re.sub(r'^#{2,3} ','',b).replace('**','') for b in blocks(method)]
        require([p.text for p in d.paragraphs]==expected,'Method body differs from authoritative source')
    else:
        body='\n'.join(''.join(p.xpath('.//w:t/text()')) for p in d._element.xpath('.//w:p'))
        require(hashlib.sha256(body.encode()).hexdigest()==BODY_HASH[label],'Authoring/example text changed from approved candidate')
        require('FSL first;' in d.tables[0].cell(0,0).text,'FSL is not first in SPEC')
    if label.startswith('Sample'):
        require('NOT RUN' in text and 'Sample edition 1.0' in text,'Sample overclaims proof or conflates versions')
        links={r.target_ref for r in d.part.rels.values() if r.is_external}
        require('https://docs.djangoproject.com/en/5.2/intro/overview/' in links,'Lost original Django reference')
def verify(root,method):
    words=method.read_text()
    require(hashlib.sha256(words.encode()).hexdigest()=='6043e0566626e403ac9509c23a43b9b7957158d2da6c1ed04a9c9a455bd2ffe7','Unexpected method revision')
    for label,(pages,_) in DOCS.items():
        stem='FSL-v1.8-'+label
        check_docx(root/(stem+'.docx'),words,label)
        with fitz.open(root/(stem+'.pdf')) as pdf:
            require(len(pdf)==pages,'Unexpected page count: '+stem)
            for page in pdf:
                text=page.get_text()
                require(AXIOM in text and 'Four-Square Template v1.8' in text,'Missing per-page footer')
                require('LOCAL CANDIDATE' not in text,'Stale PDF status')
                for x0,y0,x1,y1,*_ in page.get_text('blocks'):
                    require(x0>=0 and y0>=0 and x1<=page.rect.width+1 and y1<=page.rect.height+1,'Text outside page')
        print('PASS:',stem,'DOCX/PDF',pages,'page(s)')
def check_hashes(root):
    ledger=root/'SHA256SUMS.txt';expected={}
    for line in ledger.read_text().splitlines():
        digest,name=line.split('  ',1)
        require(name not in expected and not Path(name).is_absolute() and '..' not in Path(name).parts,'Unsafe/duplicate manifest entry')
        expected[name]=digest
    actual={p.relative_to(root).as_posix() for p in root.rglob('*') if p.is_file() and p!=ledger}
    require(actual==set(expected),'Payload membership mismatch')
    for name,digest in expected.items():
        require(hashlib.sha256((root/name).read_bytes()).hexdigest()==digest,'Payload digest mismatch: '+name)
    print('PASS: exact payload membership and',len(expected),'SHA-256 digests')
def selftest(root,method):
    label='Plain-English-Method-3P';original=root/('FSL-v1.8-'+label+'.docx')
    with tempfile.TemporaryDirectory() as td:
        bad=Path(td)/'bad.docx'
        for kind in ('underline','footer'):
            d=Document(original)
            if kind=='underline':
                for r in d._element.xpath('.//w:r'):
                    if NORM.fullmatch(''.join(r.xpath('.//w:t/text()'))):
                        r.find('w:rPr/w:u',r.nsmap).set(qn('w:val'),'none');break
            else:
                for p in d.sections[0].footer.paragraphs:
                    if AXIOM in p.text:p.text='Removed for negative test'
            d.save(bad)
            try:check_docx(bad,method.read_text(),label)
            except ValueError:print('PASS: rejected missing',kind)
            else:raise ValueError('Negative test accepted broken '+kind)
if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('root',type=Path);p.add_argument('--method',type=Path);p.add_argument('--selftest',action='store_true');p.add_argument('--hashes',action='store_true')
    a=p.parse_args();m=a.method or a.root/'FSL-v1.8-SKILL.md';verify(a.root,m)
    if a.selftest:selftest(a.root,m)
    if a.hashes:check_hashes(a.root)
