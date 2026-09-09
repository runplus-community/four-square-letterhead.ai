#!/usr/bin/env python3
"""Offline rejection checks on temporary copies; never edit the input package."""
from __future__ import annotations
import argparse
import shutil
from pathlib import Path
from tempfile import TemporaryDirectory
from docx import Document
from verify import NORM, SPECS, verify_documents, verify_policy


def expect_rejected(label, check):
    try:
        check()
    except ValueError as exc:
        print(f'PASS rejected {label}: {exc}')
    else:
        raise AssertionError(f'Mutant was accepted: {label}')


def self_test(root: Path):
    verify_policy(root)
    verify_documents(root)
    with TemporaryDirectory(prefix='fsl-v1.10-check-') as tmp:
        target=Path(tmp)/'package'
        shutil.copytree(root,target,ignore=shutil.ignore_patterns('__pycache__'))
        path=target/(SPECS[0][1]+'.docx')
        original=path.read_bytes()
        doc=Document(path)
        selected=next(r for p in doc.paragraphs for r in p.runs if NORM.search(r.text))
        selected.underline=False
        doc.save(path)
        expect_rejected('missing normative underline',lambda:verify_documents(target))
        path.write_bytes(original)
        doc=Document(path)
        for section in doc.sections:
            for p in section.footer.paragraphs:p.text=''
        doc.save(path)
        expect_rejected('missing footer',lambda:verify_documents(target))
        path.write_bytes(original)
        source=target/'FSL-v1.10-METHOD.md'
        original_source=source.read_text()
        source.write_text(original_source.replace('one start per rolling 60 minutes','unbounded starts'))
        expect_rejected('removed hourly-budget wording',lambda:verify_policy(target))
        source.write_text(original_source.replace('Intent approval, execution/spending permission and release acceptance are separate decisions', 'One approval covers everything'))
        expect_rejected('removed decision-authority separation',lambda:verify_policy(target))
        source.write_text(original_source.replace('Keep uncertain intent, authority, evidence or compatibility explicit', 'Resolve every uncertainty by guessing'))
        expect_rejected('removed uncertainty boundary',lambda:verify_policy(target))
    print('PASS 5 negative document/source checks; no live execution performed')


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--root',type=Path,default=Path(__file__).resolve().parents[1])
    self_test(p.parse_args().root.resolve())
